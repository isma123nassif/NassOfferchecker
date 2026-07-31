import csv
import hashlib
import json
import queue
import re
import shutil
import sqlite3
import subprocess
import threading
import time
import tkinter as tk
import urllib.error
import urllib.parse
import urllib.request
import warnings
from dataclasses import dataclass, asdict
from datetime import datetime
from html import unescape
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Iterable

from marketplaces import (
    chunk_products,
    build_marketplace_search_url,
    get_marketplace_config,
    get_offer_cache_path,
)


BASE_DIR = Path(__file__).resolve().parent
FASE_DIR = BASE_DIR / "fase1"
PROFILE_DIR = FASE_DIR / "browser_profile_leroy"
RUNS_DIR = FASE_DIR / "dashboard_runs"
CACHE_DB_PATH = FASE_DIR / "offer_cache.sqlite"
LOCAL_SETTINGS_PATH = BASE_DIR / "local_settings.json"
CONFORAMA_REFERENCE_MAP_PATH = FASE_DIR / "conforama_ean_mkp.xlsx"
LEROY_HOME_URL = "https://www.leroymerlin.es/"
LEROY_SEARCH_URL = "https://www.leroymerlin.es/search?q={ean}"
BLOCKED_RESOURCE_TYPES = {"image", "media", "font", "stylesheet"}
ALLOWED_EXTERNAL_HOST_PARTS = (
    "leroymerlin.es",
    "carrefour.es",
    "akamaihd.net",
    "captcha-delivery.com",
)
BLOCKED_URL_PARTS = (
    "googletagmanager",
    "google-analytics",
    "doubleclick",
    "facebook.net",
    "hotjar",
    "newrelic",
    "contentsquare",
)


@dataclass
class InputProduct:
    ean: str
    reference: str = ""
    quantity: str = ""
    price: str = ""


@dataclass
class DashboardResult:
    ean: str
    reference: str
    feed_quantity: str
    feed_price: str
    light: str
    status: str
    final_url: str
    title: str
    sku: str
    seller_name: str
    marketplace_price: str
    product_availability: str
    offer_available: bool
    buybox_ok: str
    challenge_detected: bool
    reason: str
    elapsed_seconds: float
    html_path: str


@dataclass
class WorkerConfig:
    id: int
    enabled: bool = True
    proxy_server: str = ""
    proxy_username: str = ""
    proxy_password: str = ""
    user_agent: str = ""

    @property
    def proxy_configured(self) -> bool:
        return bool(self.proxy_server.strip())

    @property
    def user_agent_configured(self) -> bool:
        return bool(self.user_agent.strip())


@dataclass
class CarrefourSearchCard:
    ean: str
    url: str
    title: str
    price: str
    seller: str
    add_to_cart: bool


@dataclass
class WortenSearchCard:
    ean: str
    url: str
    title: str
    price: str
    seller: str
    add_to_cart: bool


class ManualChallengeNotResolved(Exception):
    pass


class ConforamaApiBlocked(Exception):
    def __init__(self, status_code: int, retry_after_seconds: int = 0):
        self.status_code = status_code
        self.retry_after_seconds = retry_after_seconds
        message = f"bloqueo API Conforama HTTP {status_code}"
        if retry_after_seconds:
            message = f"{message}; retry-after {retry_after_seconds}s"
        super().__init__(message)


def read_local_settings() -> dict:
    if not LOCAL_SETTINGS_PATH.exists():
        return {}
    try:
        return json.loads(LOCAL_SETTINGS_PATH.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return {}


def bool_setting(value, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    if isinstance(value, str):
        normalized = value.strip().casefold()
        if normalized in {"1", "true", "yes", "si", "s\u00ed", "on"}:
            return True
        if normalized in {"0", "false", "no", "off"}:
            return False
    return default


def load_worker_configs(worker_count: int) -> list[WorkerConfig]:
    worker_count = max(1, min(int(worker_count), 10))
    settings = read_local_settings()
    configured: dict[int, WorkerConfig] = {}
    raw_workers = settings.get("workers", [])
    if isinstance(raw_workers, list):
        for item in raw_workers:
            if not isinstance(item, dict):
                continue
            try:
                worker_id = int(item.get("id", 0))
            except (TypeError, ValueError):
                continue
            if worker_id < 1 or worker_id > 10:
                continue
            configured[worker_id] = WorkerConfig(
                id=worker_id,
                enabled=bool(item.get("enabled", True)),
                proxy_server=str(item.get("proxy_server", "") or "").strip(),
                proxy_username=str(item.get("proxy_username", "") or "").strip(),
                proxy_password=str(item.get("proxy_password", "") or "").strip(),
                user_agent=str(item.get("user_agent", "") or "").strip(),
            )
    return [configured.get(worker_id, WorkerConfig(id=worker_id)) for worker_id in range(1, worker_count + 1)]


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip()).casefold()


def parse_quantity(value: str) -> float | None:
    value = (value or "").strip().replace(",", ".")
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def product_feed_signature(product: InputProduct) -> str:
    raw = "\n".join(
        [
            product.ean.strip(),
            product.reference.strip(),
            product.quantity.strip(),
            product.price.strip(),
        ]
    )
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def read_products_csv(path: Path) -> list[InputProduct]:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    first_line = raw.splitlines()[0] if raw.splitlines() else ""
    delimiter = ";"
    if first_line.count(",") > first_line.count(";"):
        delimiter = ","
    if "\t" in first_line and first_line.count("\t") > first_line.count(delimiter):
        delimiter = "\t"

    rows = list(csv.DictReader(raw.splitlines(), delimiter=delimiter))
    products: list[InputProduct] = []
    if rows and rows[0]:
        field_map = {normalize(name): name for name in rows[0].keys()}
        ean_key = field_map.get("ean") or field_map.get("gtin") or field_map.get("barcode")
        ref_key = field_map.get("reference") or field_map.get("sku") or field_map.get("ref")
        qty_key = field_map.get("quantity") or field_map.get("stock") or field_map.get("qty")
        price_key = field_map.get("price") or field_map.get("precio")
        if ean_key:
            for row in rows:
                ean = (row.get(ean_key) or "").strip()
                if ean:
                    products.append(
                        InputProduct(
                            ean=ean,
                            reference=(row.get(ref_key) or "").strip() if ref_key else "",
                            quantity=(row.get(qty_key) or "").strip() if qty_key else "",
                            price=(row.get(price_key) or "").strip() if price_key else "",
                        )
                    )
            return products

    for line in raw.splitlines():
        ean = line.strip().split(";")[0].strip().split(",")[0].strip()
        if ean and ean.lower() != "ean":
            products.append(InputProduct(ean=ean))
    return products


def load_conforama_reference_map(path: Path = CONFORAMA_REFERENCE_MAP_PATH) -> dict[str, str]:
    if not path.exists():
        return {}
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return {}
        normalized_headers = {normalize(str(header or "")): index for index, header in enumerate(headers)}
        ean_index = normalized_headers.get("ean")
        reference_index = (
            normalized_headers.get("sku de producto")
            or normalized_headers.get("sku producto")
            or normalized_headers.get("producto sku")
            or normalized_headers.get("reference")
            or normalized_headers.get("sku")
        )
        if ean_index is None or reference_index is None:
            return {}

        reference_map: dict[str, str] = {}
        for row in rows:
            ean = str(row[ean_index] or "").strip()
            reference = str(row[reference_index] or "").strip().upper()
            if ean.endswith(".0"):
                ean = ean[:-2]
            if reference.startswith("MKP") and ean:
                reference_map[ean] = reference
        return reference_map
    except Exception:
        return {}


def extract_candidate_urls(content: str) -> list[str]:
    patterns = [
        r"https://www\.leroymerlin\.es/productos/[^\"'<>\\\s]+?\.html",
        r"/productos/[^\"'<>\\\s]+?\.html",
    ]
    candidates: list[str] = []
    seen = set()
    for pattern in patterns:
        for match in re.findall(pattern, content):
            url = match
            if url.startswith("/"):
                url = "https://www.leroymerlin.es" + url
            url = url.split("?")[0]
            if url not in seen:
                seen.add(url)
                candidates.append(url)
    return candidates


def extract_product_jsonld(content: str) -> dict:
    scripts = re.findall(
        r'<script[^>]+type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
        content,
        flags=re.IGNORECASE | re.DOTALL,
    )
    for script in scripts:
        raw = unescape(script).strip()
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            continue
        entries = data if isinstance(data, list) else [data]
        for entry in entries:
            if isinstance(entry, dict) and entry.get("@type") == "Product":
                return entry
    return {}


def visible_text(content: str) -> str:
    text = re.sub(r"<script\b[^>]*>.*?</script>", " ", content, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"<style\b[^>]*>.*?</style>", " ", text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", unescape(text)).strip().lower()


def challenge_detected(content: str, title: str, url: str) -> bool:
    lowered = f"{title}\n{url}\n{content[:200000]}".lower()
    text = visible_text(content[:200000])
    usable = len(content.encode("utf-8")) > 50_000 and "leroy merlin" in lowered
    has_product_signal = (
        'id="jsonld_product"' in lowered
        or '"add_to_cart_availability":true' in content
        or "schema.org/discontinued" in lowered
        or "recommendation-no-offer-banner" in lowered
    )
    has_home_signal = (
        "bricolaje, decoraci" in text
        and "leroy merlin" in text
        and "mi carrito" in text
        and "productos" in text
    )
    if usable and has_product_signal and "please enable js" not in lowered:
        return False
    if usable and has_home_signal and "var dd=" not in lowered and "var dd =" not in lowered:
        return False
    if "captcha-delivery.com" in url.lower():
        return True
    visible_markers = [
        "please enable js",
        "access denied",
        "recaptcha",
        "g-recaptcha",
        "hcaptcha",
        "are you human",
        "eres humano",
        "verifica que eres",
        "unusual traffic",
        "blocked",
        "bloqueado",
        "bot detection",
    ]
    structural_markers = ["var dd=", "var dd =", "g-recaptcha", "hcaptcha"]
    return any(marker in text for marker in visible_markers) or any(marker in lowered for marker in structural_markers)


def extract_availability(content: str) -> tuple[str, bool, str]:
    product = extract_product_jsonld(content)
    offers = product.get("offers", {}) if isinstance(product, dict) else {}
    availability = ""
    if isinstance(offers, dict):
        availability = str(offers.get("availability", "") or "")

    availability_lower = availability.lower()
    if "discontinued" in availability_lower or "outofstock" in availability_lower:
        return availability or "negative_marker", False, f"jsonld availability={availability}"
    if '"add_to_cart_availability":true' in content:
        return availability or "add_to_cart_available", True, "add_to_cart_availability=true"
    if "recommendation-no-offer-banner" in content:
        return availability or "recommendation_no_offer_banner", False, "recommendation-no-offer-banner"
    if '"total_offer_count":0' in content and '"name":"cdl_products"' in content:
        return availability or "zero_offer_count", False, "cdl_products total_offer_count=0"

    text = visible_text(content)
    for marker in ("este producto no está disponible", "este producto no esta disponible", "producto no disponible"):
        if marker in text:
            return availability or "text_unavailable", False, "texto de no disponibilidad"
    return availability or "unknown", False, "sin senal comprable"


def extract_first(pattern: str, content: str) -> str:
    match = re.search(pattern, content, flags=re.IGNORECASE | re.DOTALL)
    return unescape(match.group(1)).strip() if match else ""


def extract_product_fields(content: str) -> dict[str, str]:
    product = extract_product_jsonld(content)
    offers = product.get("offers", {}) if isinstance(product, dict) else {}
    sku = extract_first(r'"sku"\s*:\s*"([^"]+)"', content)
    seller_name = extract_first(r'"seller_name"\s*:\s*"([^"]+)"', content)
    price = extract_first(r'"unitprice_ati"\s*:\s*([0-9]+(?:\.[0-9]+)?)', content)
    if not price and isinstance(offers, dict):
        price = str(offers.get("price", "") or "")
    return {"sku": sku, "seller_name": seller_name, "marketplace_price": price}


def classify_light(
    product: InputProduct,
    content: str,
    title: str,
    final_url: str,
    expected_seller: str,
) -> tuple[str, str, str, bool, str, str, str, str]:
    candidates = extract_candidate_urls(content)
    is_challenge = challenge_detected(content, title, final_url)
    availability, offer_available, availability_evidence = extract_availability(content)
    fields = extract_product_fields(content)
    seller_name = fields["seller_name"]
    feed_qty = parse_quantity(product.quantity)

    if is_challenge:
        return "gray", "INCIERTA", availability, offer_available, "unknown", seller_name, fields["marketplace_price"], "challenge detectado"
    if feed_qty is not None and feed_qty <= 0:
        return "red", "SIN_STOCK_FEED", availability, offer_available, "n/a", seller_name, fields["marketplace_price"], "stock feed <= 0"
    if len(content.encode("utf-8")) < 10_000:
        return "gray", "INCIERTA", availability, offer_available, "unknown", seller_name, fields["marketplace_price"], "HTML demasiado pequeno"
    if product.ean not in content and not candidates:
        return "red", "NO_VIVA", availability, offer_available, "n/a", seller_name, fields["marketplace_price"], "no se vende en marketplace"
    if product.ean not in content and candidates:
        return "gray", "INCIERTA", availability, offer_available, "unknown", seller_name, fields["marketplace_price"], "candidatos sin EAN confirmado"
    if not offer_available:
        seller_name = ""
        fields["marketplace_price"] = ""
        if feed_qty is not None and feed_qty > 0:
            return "yellow", "NO_DISPONIBLE_CON_STOCK", availability, offer_available, "n/a", seller_name, fields["marketplace_price"], availability_evidence
        return "red", "NO_VIVA", availability, offer_available, "n/a", seller_name, fields["marketplace_price"], availability_evidence

    buybox_ok = "n/a"
    if expected_seller.strip():
        buybox_ok = "yes" if normalize(seller_name) == normalize(expected_seller) else "no"
        if buybox_ok == "no":
            return "yellow", "BUYBOX_PERDIDA", availability, offer_available, buybox_ok, seller_name, fields["marketplace_price"], f"seller actual={seller_name}"

    return "green", "OK", availability, offer_available, buybox_ok, seller_name, fields["marketplace_price"], "oferta disponible"


def write_dashboard_row(path: Path, row: DashboardResult) -> None:
    exists = path.exists()
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(asdict(row).keys()))
        if not exists:
            writer.writeheader()
        writer.writerow(asdict(row))


class OfferCache:
    def __init__(self, path: Path = CACHE_DB_PATH):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._init_db()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.path)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_db(self) -> None:
        with self._connect() as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS offer_results (
                    ean TEXT NOT NULL,
                    expected_seller TEXT NOT NULL,
                    feed_signature TEXT NOT NULL,
                    checked_at REAL NOT NULL,
                    result_json TEXT NOT NULL,
                    PRIMARY KEY (ean, expected_seller)
                )
                """
            )

    def get_fresh(
        self,
        product: InputProduct,
        expected_seller: str,
        ttl_hours: float,
    ) -> DashboardResult | None:
        expected_seller = (expected_seller or "").strip()
        signature = product_feed_signature(product)
        with self._connect() as conn:
            row = conn.execute(
                """
                SELECT feed_signature, checked_at, result_json
                FROM offer_results
                WHERE ean = ? AND expected_seller = ?
                """,
                (product.ean, expected_seller),
            ).fetchone()
        if not row:
            return None
        if row["feed_signature"] != signature:
            return None
        max_age_seconds = max(ttl_hours, 0) * 3600
        if max_age_seconds and (time.time() - float(row["checked_at"])) > max_age_seconds:
            return None
        data = json.loads(row["result_json"])
        if data.get("status") == "INCIERTA" or data.get("light") == "gray":
            return None
        data.update(
            {
                "reference": product.reference,
                "feed_quantity": product.quantity,
                "feed_price": product.price,
                "reason": f"cache diferencial: {data.get('reason', '')}".strip(),
                "elapsed_seconds": 0.0,
            }
        )
        return DashboardResult(**data)

    def put(self, product: InputProduct, expected_seller: str, row: DashboardResult) -> None:
        if row.status == "INCIERTA" or row.light == "gray":
            return
        data = asdict(row)
        with self._connect() as conn:
            conn.execute(
                """
                INSERT INTO offer_results (ean, expected_seller, feed_signature, checked_at, result_json)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(ean, expected_seller) DO UPDATE SET
                    feed_signature = excluded.feed_signature,
                    checked_at = excluded.checked_at,
                    result_json = excluded.result_json
                """,
                (
                    product.ean,
                    (expected_seller or "").strip(),
                    product_feed_signature(product),
                    time.time(),
                    json.dumps(data, ensure_ascii=False),
                ),
            )


class AlertNotifier:
    def __init__(self):
        self.settings = self._load_settings()
        self.enabled = self._setting_bool("alerts_enabled", True)
        self.slack_webhook_url = (
            self._env("SLACK_WEBHOOK_URL")
            or str(self.settings.get("slack_webhook_url", "") or "").strip()
        )
        self.cooldown_seconds = int(self._setting_float("alert_cooldown_minutes", 15.0) * 60)
        self.sent_at: dict[str, float] = {}

    def _load_settings(self) -> dict:
        return read_local_settings()

    def _env(self, name: str) -> str:
        import os

        return os.environ.get(name, "").strip()

    def _setting_bool(self, name: str, default: bool) -> bool:
        value = self._env(name.upper())
        if not value:
            value = self.settings.get(name, default)
        if isinstance(value, bool):
            return value
        return str(value).strip().lower() not in {"0", "false", "no", "off"}

    def _setting_float(self, name: str, default: float) -> float:
        value = self._env(name.upper())
        if not value:
            value = self.settings.get(name, default)
        try:
            return float(str(value).replace(",", "."))
        except ValueError:
            return default

    def should_alert(self, row: DashboardResult) -> bool:
        evidence = f"{row.status} {row.reason} {row.title} {row.final_url}".casefold()
        markers = (
            "captcha",
            "recaptcha",
            "hcaptcha",
            "challenge",
            "datadome",
            "access denied",
            "blocked",
            "bloqueo",
            "error de navegacion",
            "timeout",
            "please enable js",
        )
        return row.challenge_detected or row.status == "INCIERTA" or any(marker in evidence for marker in markers)

    def notify_if_needed(self, row: DashboardResult, run_dir: Path, index: int, total: int) -> bool:
        if not self.enabled or not self.slack_webhook_url or not self.should_alert(row):
            return False
        key = row.status if row.status else "technical_alert"
        now = time.time()
        if self.cooldown_seconds > 0 and now - self.sent_at.get(key, 0) < self.cooldown_seconds:
            return False
        self.sent_at[key] = now
        return self.send_slack(row, run_dir, index, total)

    def send_slack(self, row: DashboardResult, run_dir: Path, index: int, total: int) -> bool:
        payload = {
            "text": f"OfferChecker alerta: {row.status} en EAN {row.ean}",
            "blocks": [
                {
                    "type": "section",
                    "text": {
                        "type": "mrkdwn",
                        "text": f"*OfferChecker alerta*\nDetectado posible bloqueo/challenge en marketplace.",
                    },
                },
                {
                    "type": "section",
                    "fields": [
                        {"type": "mrkdwn", "text": f"*EAN:*\n`{row.ean}`"},
                        {"type": "mrkdwn", "text": f"*Estado:*\n`{row.status}`"},
                        {"type": "mrkdwn", "text": f"*Progreso:*\n{index}/{total}"},
                        {"type": "mrkdwn", "text": f"*Challenge:*\n{row.challenge_detected}"},
                    ],
                },
                {
                    "type": "section",
                    "text": {"type": "mrkdwn", "text": f"*Motivo:*\n{row.reason[:600]}"},
                },
                {
                    "type": "context",
                    "elements": [
                        {"type": "mrkdwn", "text": f"Run local: `{run_dir}`"},
                        {"type": "mrkdwn", "text": f"URL: {row.final_url}"},
                    ],
                },
            ],
        }
        data = json.dumps(payload).encode("utf-8")
        request = urllib.request.Request(
            self.slack_webhook_url,
            data=data,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=10) as response:
                return 200 <= response.status < 300
        except (OSError, urllib.error.URLError):
            return False


def minimize_chromium_window(context, page) -> None:
    try:
        session = context.new_cdp_session(page)
        window = session.send("Browser.getWindowForTarget")
        window_id = window.get("windowId")
        if window_id is not None:
            session.send(
                "Browser.setWindowBounds",
                {"windowId": window_id, "bounds": {"windowState": "minimized"}},
            )
    except Exception:
        pass


def show_chromium_window(context, page) -> None:
    try:
        page.bring_to_front()
    except Exception:
        pass
    try:
        session = context.new_cdp_session(page)
        window = session.send("Browser.getWindowForTarget")
        window_id = window.get("windowId")
        if window_id is not None:
            session.send(
                "Browser.setWindowBounds",
                {"windowId": window_id, "bounds": {"windowState": "maximized"}},
            )
    except Exception:
        pass


def install_lightweight_routes(context, minimal_data_mode: bool = True) -> None:
    def handle_route(route) -> None:
        try:
            request = route.request
            url = request.url.lower()
            host = ""
            match = re.match(r"^https?://([^/]+)", url)
            if match:
                host = match.group(1)
            if request.resource_type in BLOCKED_RESOURCE_TYPES or any(part in url for part in BLOCKED_URL_PARTS):
                route.abort()
            elif minimal_data_mode and host and not any(part in host for part in ALLOWED_EXTERNAL_HOST_PARTS):
                route.abort()
            else:
                route.continue_()
        except Exception:
            pass

    context.route("**/*", handle_route)


class LeroyChecker:
    marketplace_key = "leroy"

    def __init__(
        self,
        products: list[InputProduct],
        expected_seller: str,
        delay: float,
        nav_timeout: float = 18.0,
        retry_timeout: float = 35.0,
        block_assets: bool = True,
        minimal_data_mode: bool = True,
        data_settle_ms: int = 700,
        worker_count: int = 1,
        circuit_breaker_threshold: int = 4,
        cache_path: Path | None = None,
    ):
        self.marketplace_config = get_marketplace_config(self.marketplace_key)
        self.products = products
        self.expected_seller = expected_seller
        self.delay = delay
        self.nav_timeout_ms = int(max(nav_timeout, 5.0) * 1000)
        self.retry_timeout_ms = int(max(retry_timeout, nav_timeout, 5.0) * 1000)
        self.block_assets = block_assets
        self.minimal_data_mode = minimal_data_mode
        self.data_settle_ms = max(0, min(int(data_settle_ms), 3000))
        self.worker_configs = load_worker_configs(worker_count)
        self.active_worker_configs = [config for config in self.worker_configs if config.enabled]
        if not self.active_worker_configs:
            self.active_worker_configs = [WorkerConfig(id=1)]
        self.worker_count = len(self.active_worker_configs)
        self.circuit_breaker_threshold = max(1, int(circuit_breaker_threshold))
        self.stop_requested = False
        self.circuit_open = False
        self.completed_count = 0
        self.challenge_streak = 0
        self.result_lock = threading.Lock()
        self.run_dir = RUNS_DIR / datetime.now().strftime("%Y%m%d_%H%M%S")
        self.summary_path = self.run_dir / "summary_live.csv"
        self.cache = OfferCache(cache_path or get_offer_cache_path(self.marketplace_key))
        self.alerts = AlertNotifier()
        self.events: queue.Queue | None = None

    def stop(self) -> None:
        self.stop_requested = True

    def profile_dir_for_worker(self, worker_id: int) -> Path:
        if worker_id <= 1:
            return PROFILE_DIR
        return FASE_DIR / f"browser_profile_leroy_worker_{worker_id}"

    def launch_worker_context(self, p, worker_id: int, PlaywrightTimeoutError):
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(self.profile_dir_for_worker(worker_id)),
            headless=False,
            viewport={"width": 1365, "height": 900},
            locale="es-ES",
            timezone_id="Europe/Madrid",
            args=["--disable-blink-features=AutomationControlled", "--start-minimized"],
        )
        if self.block_assets:
            install_lightweight_routes(context, self.minimal_data_mode)
        page = context.pages[0] if context.pages else context.new_page()
        minimize_chromium_window(context, page)
        page.set_default_timeout(self.retry_timeout_ms)
        warmup_error = ""
        try:
            page.goto(self.marketplace_config.home_url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
        except PlaywrightTimeoutError as exc:
            warmup_error = f"timeout warm-up: {exc}"
        self.settle_and_stop(page)
        try:
            content = page.content()
            title = page.title()
            final_url = page.url
        except Exception as exc:
            return context, page, False, f"error warm-up: {exc}"
        if challenge_detected(content, title, final_url):
            html_file = self.run_dir / f"worker_{worker_id:02d}_warmup_challenge.html"
            html_file.write_text(content, encoding="utf-8")
            return context, page, False, "challenge en warm-up"
        return context, page, True, warmup_error or "warm-up correcto"

    def notify_worker_blocked(
        self,
        worker_id: int,
        reason: str,
        page,
        events: queue.Queue,
    ) -> None:
        try:
            final_url = page.url
            title = page.title()
        except Exception:
            final_url = self.marketplace_config.home_url
            title = ""
        row = DashboardResult(
            ean=f"WORKER_{worker_id}_WARMUP",
            reference="",
            feed_quantity="",
            feed_price="",
            light="gray",
            status="INCIERTA",
            final_url=final_url,
            title=title,
            sku="",
            seller_name="",
            marketplace_price="",
            product_availability="unknown",
            offer_available=False,
            buybox_ok="unknown",
            challenge_detected=True,
            reason=reason,
            elapsed_seconds=0,
            html_path=str(self.run_dir / f"worker_{worker_id:02d}_warmup_challenge.html"),
        )
        if self.alerts.notify_if_needed(row, self.run_dir, 0, len(self.products)):
            events.put(("alert_sent", row.ean, row.status))

    def settle_and_stop(self, page) -> None:
        if self.data_settle_ms:
            try:
                page.wait_for_timeout(self.data_settle_ms)
            except Exception:
                pass
        if self.minimal_data_mode:
            try:
                page.evaluate("window.stop()")
            except Exception:
                pass

    def check_one_product(self, context, page, product: InputProduct, index: int, PlaywrightError) -> tuple[DashboardResult, object]:
        item_started = time.monotonic()
        url = build_marketplace_search_url(self.marketplace_key, [product.ean])
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
        except PlaywrightError:
            try:
                page.close()
            except PlaywrightError:
                pass
            page = context.new_page()
            minimize_chromium_window(context, page)
            page.set_default_timeout(self.retry_timeout_ms)
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=self.retry_timeout_ms)
            except PlaywrightError as exc:
                row = DashboardResult(
                    ean=product.ean,
                    reference=product.reference,
                    feed_quantity=product.quantity,
                    feed_price=product.price,
                    light="yellow",
                    status="INCIERTA",
                    final_url=url,
                    title="",
                    sku="",
                    seller_name="",
                    marketplace_price="",
                    product_availability="unknown",
                    offer_available=False,
                    buybox_ok="unknown",
                    challenge_detected=False,
                    reason=f"error de navegacion: {exc}",
                    elapsed_seconds=round(time.monotonic() - item_started, 2),
                    html_path="",
                )
                return row, page

        self.settle_and_stop(page)
        content = page.content()
        title = page.title()
        final_url = page.url
        light, status, availability, offer_available, buybox_ok, seller_name, marketplace_price, reason = classify_light(
            product, content, title, final_url, self.expected_seller
        )
        html_path = ""
        if light != "green":
            html_file = self.run_dir / f"{index:04d}_{product.ean}.html"
            html_file.write_text(content, encoding="utf-8")
            html_path = str(html_file)

        row = DashboardResult(
            ean=product.ean,
            reference=product.reference,
            feed_quantity=product.quantity,
            feed_price=product.price,
            light=light,
            status=status,
            final_url=final_url,
            title=title,
            sku=extract_product_fields(content)["sku"],
            seller_name=seller_name,
            marketplace_price=marketplace_price,
            product_availability=availability,
            offer_available=offer_available,
            buybox_ok=buybox_ok,
            challenge_detected=challenge_detected(content, title, final_url),
            reason=reason,
            elapsed_seconds=round(time.monotonic() - item_started, 2),
            html_path=html_path,
        )
        return row, page

    def record_result(self, row: DashboardResult, index: int, total: int, events: queue.Queue) -> int:
        with self.result_lock:
            write_dashboard_row(self.summary_path, row)
            self.cache.put(InputProduct(row.ean, row.reference, row.feed_quantity, row.feed_price), self.expected_seller, row)
            if self.alerts.notify_if_needed(row, self.run_dir, index, total):
                events.put(("alert_sent", row.ean, row.status))
            if self.is_circuit_signal(row):
                self.challenge_streak += 1
            else:
                self.challenge_streak = 0
            if self.challenge_streak >= self.circuit_breaker_threshold and not self.circuit_open:
                self.circuit_open = True
                self.stop_requested = True
                events.put(("circuit_breaker", row.ean, row.status, self.challenge_streak))
            self.completed_count += 1
            return self.completed_count

    def is_circuit_signal(self, row: DashboardResult) -> bool:
        return self.alerts.should_alert(row)

    def browser_worker(
        self,
        worker_id: int,
        work_queue: queue.Queue,
        events: queue.Queue,
        PlaywrightError,
        PlaywrightTimeoutError,
        sync_playwright,
    ) -> None:
        context = None
        final_status = "cerrado"
        try:
            with sync_playwright() as p:
                events.put(("worker_status", worker_id, "calentando"))
                context, page, healthy, warmup_reason = self.launch_worker_context(p, worker_id, PlaywrightTimeoutError)
                if not healthy:
                    final_status = "bloqueado"
                    events.put(("worker_status", worker_id, "bloqueado"))
                    events.put(("worker_blocked", worker_id, warmup_reason))
                    self.notify_worker_blocked(worker_id, warmup_reason, page, events)
                    return
                events.put(("worker_status", worker_id, "activo"))
                total = len(self.products)
                while not self.stop_requested:
                    try:
                        index, product = work_queue.get_nowait()
                    except queue.Empty:
                        break
                    try:
                        if self.stop_requested:
                            break
                        row, page = self.check_one_product(context, page, product, index, PlaywrightError)
                        completed = self.record_result(row, index, total, events)
                        events.put(("result", completed, total, row))
                        if not self.stop_requested and completed < total:
                            time.sleep(self.delay)
                    finally:
                        work_queue.task_done()
        except Exception as exc:
            self.stop_requested = True
            events.put(("error", f"worker {worker_id}: {exc}"))
        finally:
            cdp_browsers = getattr(self, "cdp_browsers", {})
            cdp_processes = getattr(self, "cdp_processes", {})
            browser = cdp_browsers.pop(worker_id, None)
            if browser is not None:
                try:
                    browser.close()
                except Exception:
                    pass
                process = cdp_processes.pop(worker_id, None)
                if process is not None and process.poll() is None:
                    try:
                        process.terminate()
                    except Exception:
                        pass
            elif context is not None:
                try:
                    context.close()
                except Exception:
                    pass
            process = cdp_processes.pop(worker_id, None)
            if process is not None and process.poll() is None:
                try:
                    process.terminate()
                except Exception:
                    pass
            if final_status != "bloqueado":
                events.put(("worker_status", worker_id, final_status))

    def run(self, events: queue.Queue) -> None:
        self.events = events
        started = time.monotonic()
        try:
            from playwright.sync_api import Error as PlaywrightError
            from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            events.put(("error", f"Playwright no esta instalado: {exc}"))
            return

        self.run_dir.mkdir(parents=True, exist_ok=True)
        events.put(("run_dir", str(self.run_dir)))
        events.put(("phase3", self.worker_count, self.circuit_breaker_threshold))
        for config in self.worker_configs:
            events.put(
                (
                    "worker_config",
                    config.id,
                    config.enabled,
                    config.proxy_configured,
                    config.user_agent_configured,
                )
            )
            if not config.enabled:
                events.put(("worker_status", config.id, "desactivado"))

        work_queue: queue.Queue = queue.Queue()
        for index, product in enumerate(self.products, start=1):
            work_queue.put((index, product))

        worker_threads = []
        for config in self.active_worker_configs:
            thread = threading.Thread(
                target=self.browser_worker,
                args=(config.id, work_queue, events, PlaywrightError, PlaywrightTimeoutError, sync_playwright),
                daemon=True,
            )
            worker_threads.append(thread)
            thread.start()

        for thread in worker_threads:
            thread.join()

        elapsed = round(time.monotonic() - started, 1)
        summary = str(self.summary_path)
        if self.circuit_open:
            summary = f"{summary} - corte tecnico activo"
        events.put(("done", elapsed, summary))


class CarrefourChecker(LeroyChecker):
    marketplace_key = "carrefour"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.block_assets = False
        self.minimal_data_mode = False
        self.data_settle_ms = max(self.data_settle_ms, 2500)
        self.cdp_browsers = {}
        self.cdp_processes = {}

    def profile_dir_for_worker(self, worker_id: int) -> Path:
        return FASE_DIR / f"browser_profile_carrefour_worker_{worker_id}"

    def launch_worker_context(self, p, worker_id: int, PlaywrightTimeoutError):
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(self.profile_dir_for_worker(worker_id)),
            headless=False,
            viewport={"width": 1365, "height": 900},
            locale="es-ES",
            timezone_id="Europe/Madrid",
            args=["--disable-blink-features=AutomationControlled", "--disable-quic", "--start-minimized"],
        )
        page = context.pages[0] if context.pages else context.new_page()
        minimize_chromium_window(context, page)
        page.set_default_timeout(self.retry_timeout_ms)
        warmup_error = ""
        loaded = False
        for attempt, timeout_ms in enumerate((self.nav_timeout_ms, self.retry_timeout_ms), start=1):
            try:
                page.goto(self.marketplace_config.home_url, wait_until="domcontentloaded", timeout=timeout_ms)
                loaded = True
                break
            except PlaywrightTimeoutError as exc:
                warmup_error = f"timeout warm-up Carrefour intento {attempt}: {exc}"
            except Exception as exc:
                warmup_error = f"error warm-up Carrefour intento {attempt}: {exc}"
            try:
                page.wait_for_timeout(1000)
            except Exception:
                pass
        if not loaded:
            return context, page, False, warmup_error or "error warm-up Carrefour"

        self.settle_and_stop(page)
        try:
            content = page.content()
            title = page.title()
            final_url = page.url
        except Exception as exc:
            return context, page, False, f"error warm-up Carrefour: {exc}"
        if challenge_detected(content, title, final_url):
            html_file = self.run_dir / f"worker_{worker_id:02d}_warmup_challenge.html"
            html_file.write_text(content, encoding="utf-8")
            return context, page, False, "challenge en warm-up Carrefour"
        return context, page, True, warmup_error or "warm-up Carrefour correcto"

    def settle_and_stop(self, page) -> None:
        try:
            page.wait_for_load_state("load", timeout=min(self.retry_timeout_ms, 12000))
        except Exception:
            pass
        try:
            page.wait_for_load_state("networkidle", timeout=min(self.retry_timeout_ms, 8000))
        except Exception:
            pass
        try:
            page.wait_for_timeout(self.data_settle_ms)
        except Exception:
            pass

    def extract_search_cards(self, content: str) -> list[CarrefourSearchCard]:
        cards: list[CarrefourSearchCard] = []
        articles = re.findall(
            r"<article\b(?=[^>]*data-test=[\"']search-grid-result[\"'])[\s\S]*?</article>",
            content,
            flags=re.IGNORECASE,
        )
        for article in articles:
            ean = extract_first(r'data-scroll=["\'](\d{8,14})["\']', article)
            url = extract_first(r'href=["\'](https://www\.carrefour\.es/[^"\']+/(\d{8,14})/p)["\']', article)
            if not ean:
                ean = extract_first(r"https://www\.carrefour\.es/[^\"']+/(\d{8,14})/p", article)
            if not ean:
                ean = extract_first(r"/(\d{8,14})_1\.(?:jpg|png|webp)", article)

            price = extract_first(
                r'data-test=["\']result-current-price["\'][\s\S]*?<span[^>]*class=["\'][^"\']*x-currency[^"\']*["\'][^>]*>(.*?)</span>',
                article,
            )
            seller = extract_first(r"Vendido por\s*<span[^>]*>(.*?)</span>", article)
            title_text = extract_first(
                r'data-test=["\']result-title["\'][\s\S]*?<p[^>]*>(.*?)</p>',
                article,
            )
            if not title_text:
                title_text = extract_first(r'<img[^>]+alt=["\']([^"\']+)["\']', article)
            add_to_cart = 'data-test="result-add-to-cart"' in article or "data-test='result-add-to-cart'" in article

            if ean or price or seller or title_text:
                cards.append(
                    CarrefourSearchCard(
                        ean=ean,
                        url=url,
                        title=title_text,
                        price=price,
                        seller=seller,
                        add_to_cart=add_to_cart,
                    )
                )
        return cards

    def no_exact_match_for_ean(self, content_lower: str, text: str, ean: str) -> bool:
        if not ean or ean not in text:
            return False
        if 'data-test="spellcheck-message"' not in content_lower and "no hemos encontrado coincidencias" not in text:
            return False
        patterns = (
            rf"no hemos encontrado coincidencias para\s+{re.escape(ean)}\b",
            rf"no se han encontrado coincidencias para\s+{re.escape(ean)}\b",
        )
        return any(re.search(pattern, text) for pattern in patterns)

    def classify_batch(
        self,
        batch: list[InputProduct],
        content: str,
        title: str,
        final_url: str,
        elapsed_seconds: float,
        html_path: str,
    ) -> list[DashboardResult]:
        is_challenge = challenge_detected(content, title, final_url)
        content_lower = content.lower()
        text = visible_text(content)
        html_too_small = len(content.encode("utf-8")) < 10_000
        no_results_signal = any(
            marker in text
            for marker in (
                "no se han encontrado resultados",
                "no hemos encontrado resultados",
                "sin resultados",
                "ningun resultado",
                "ningún resultado",
            )
        )
        product_signal = any(
            marker in content_lower
            for marker in (
                "product-card",
                "search-grid-result",
                "producttile",
                "product_tile",
                "data-product",
                "product-list",
                "productgrid",
                "add-to-cart",
                '"price"',
                '"offers"',
            )
        )
        cards = self.extract_search_cards(content)
        cards_by_ean = {card.ean: card for card in cards if card.ean}
        rows: list[DashboardResult] = []
        for position, product in enumerate(batch):
            no_exact_match = self.no_exact_match_for_ean(content_lower, text, product.ean)
            card = cards_by_ean.get(product.ean)
            if not card and not no_exact_match and len(cards) == len(batch) and position < len(cards):
                card = cards[position]

            if is_challenge:
                light = "gray"
                status = "INCIERTA"
                reason = "challenge detectado"
                seller_name = ""
                marketplace_price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"
            elif html_too_small:
                light = "gray"
                status = "INCIERTA"
                reason = "HTML demasiado pequeno"
                seller_name = ""
                marketplace_price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"
            elif no_exact_match:
                light = "red"
                status = "NO_VIVA"
                reason = "Carrefour no encontro coincidencia exacta para el EAN"
                seller_name = ""
                marketplace_price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "no_exact_match"
            elif card:
                seller_name = card.seller
                marketplace_price = card.price
                result_url = card.url or final_url
                product_title = card.title or title
                offer_available = card.add_to_cart
                availability = "search_result_available" if card.add_to_cart else "search_result_without_add_button"
                buybox_ok = "unknown"
                if self.expected_seller.strip() and seller_name:
                    seller_compact = re.sub(r"[^a-z0-9]+", "", normalize(seller_name))
                    expected_compact = re.sub(r"[^a-z0-9]+", "", normalize(self.expected_seller))
                    buybox_ok = "yes" if seller_compact == expected_compact else "no"

                if not card.add_to_cart:
                    light = "yellow"
                    status = "NO_DISPONIBLE_CON_STOCK"
                    reason = "tarjeta Carrefour encontrada sin boton de compra"
                elif buybox_ok == "no":
                    light = "yellow"
                    status = "BUYBOX_PERDIDA"
                    reason = f"seller actual={seller_name}"
                elif self.expected_seller.strip() and not seller_name:
                    light = "yellow"
                    status = "ENCONTRADO_PENDIENTE_DETALLE"
                    reason = "tarjeta Carrefour encontrada; seller no extraido"
                else:
                    light = "green"
                    status = "OK"
                    reason = "oferta Carrefour encontrada desde busqueda"
            elif no_results_signal:
                light = "red"
                status = "NO_VIVA"
                reason = "EAN no encontrado en busqueda Carrefour"
                seller_name = ""
                marketplace_price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "no_results"
            elif product_signal:
                light = "gray"
                status = "INCIERTA"
                reason = "Carrefour devolvio resultados, pero no se pudo mapear tarjeta al EAN"
                seller_name = ""
                marketplace_price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"
            else:
                light = "gray"
                status = "INCIERTA"
                reason = "Carrefour sin senal clara de resultados ni de no resultados"
                seller_name = ""
                marketplace_price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"

            rows.append(
                DashboardResult(
                    ean=product.ean,
                    reference=product.reference,
                    feed_quantity=product.quantity,
                    feed_price=product.price,
                    light=light,
                    status=status,
                    final_url=result_url,
                    title=product_title,
                    sku=card.ean if card else "",
                    seller_name=seller_name,
                    marketplace_price=marketplace_price,
                    product_availability=availability,
                    offer_available=offer_available,
                    buybox_ok=buybox_ok,
                    challenge_detected=is_challenge,
                    reason=reason,
                    elapsed_seconds=elapsed_seconds,
                    html_path=html_path,
                )
            )
        return rows

    def check_product_batch(
        self,
        context,
        page,
        batch: list[InputProduct],
        batch_index: int,
        PlaywrightError,
    ) -> tuple[list[DashboardResult], object]:
        item_started = time.monotonic()
        url = build_marketplace_search_url(self.marketplace_key, [product.ean for product in batch])
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
        except PlaywrightError:
            try:
                page.close()
            except PlaywrightError:
                pass
            page = context.new_page()
            minimize_chromium_window(context, page)
            page.set_default_timeout(self.retry_timeout_ms)
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=self.retry_timeout_ms)
            except PlaywrightError as exc:
                elapsed = round(time.monotonic() - item_started, 2)
                rows = []
                for product in batch:
                    rows.append(
                        DashboardResult(
                            ean=product.ean,
                            reference=product.reference,
                            feed_quantity=product.quantity,
                            feed_price=product.price,
                            light="gray",
                            status="INCIERTA",
                            final_url=url,
                            title="",
                            sku="",
                            seller_name="",
                            marketplace_price="",
                            product_availability="unknown",
                            offer_available=False,
                            buybox_ok="unknown",
                            challenge_detected=False,
                            reason=f"error de navegacion Carrefour batch: {exc}",
                            elapsed_seconds=elapsed,
                            html_path="",
                        )
                    )
                return rows, page

        self.settle_and_stop(page)
        content = page.content()
        title = page.title()
        final_url = page.url
        elapsed = round(time.monotonic() - item_started, 2)
        html_file = self.run_dir / f"carrefour_batch_{batch_index:04d}.html"
        html_file.write_text(content, encoding="utf-8")
        rows = self.classify_batch(batch, content, title, final_url, elapsed, str(html_file))
        rows, page = self.retry_unmapped_batch_rows(context, page, rows, batch, batch_index, PlaywrightError)
        return rows, page

    def retry_unmapped_batch_rows(
        self,
        context,
        page,
        rows: list[DashboardResult],
        batch: list[InputProduct],
        batch_index: int,
        PlaywrightError,
    ) -> tuple[list[DashboardResult], object]:
        products_by_ean = {product.ean: product for product in batch}
        updated_rows = list(rows)
        retry_reason = "Carrefour devolvio resultados, pero no se pudo mapear tarjeta al EAN"
        for position, row in enumerate(rows, start=1):
            if row.status != "INCIERTA" or row.reason != retry_reason:
                continue
            product = products_by_ean.get(row.ean)
            if not product or self.stop_requested:
                continue
            retry_started = time.monotonic()
            url = build_marketplace_search_url(self.marketplace_key, [product.ean])
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
            except PlaywrightError:
                try:
                    page.close()
                except PlaywrightError:
                    pass
                page = context.new_page()
                minimize_chromium_window(context, page)
                page.set_default_timeout(self.retry_timeout_ms)
                try:
                    page.goto(url, wait_until="domcontentloaded", timeout=self.retry_timeout_ms)
                except PlaywrightError:
                    continue

            self.settle_and_stop(page)
            content = page.content()
            title = page.title()
            final_url = page.url
            elapsed = round(time.monotonic() - retry_started, 2)
            html_file = self.run_dir / f"carrefour_batch_{batch_index:04d}_retry_{position:02d}_{product.ean}.html"
            html_file.write_text(content, encoding="utf-8")
            retry_rows = self.classify_batch([product], content, title, final_url, elapsed, str(html_file))
            if retry_rows:
                updated_rows[position - 1] = retry_rows[0]
        return updated_rows, page

    def browser_worker(
        self,
        worker_id: int,
        work_queue: queue.Queue,
        events: queue.Queue,
        PlaywrightError,
        PlaywrightTimeoutError,
        sync_playwright,
    ) -> None:
        context = None
        final_status = "cerrado"
        try:
            with sync_playwright() as p:
                events.put(("worker_status", worker_id, "calentando"))
                context, page, healthy, warmup_reason = self.launch_worker_context(p, worker_id, PlaywrightTimeoutError)
                if not healthy:
                    final_status = "bloqueado"
                    events.put(("worker_status", worker_id, "bloqueado"))
                    events.put(("worker_blocked", worker_id, warmup_reason))
                    self.notify_worker_blocked(worker_id, warmup_reason, page, events)
                    return
                events.put(("worker_status", worker_id, "activo"))
                total = len(self.products)
                while not self.stop_requested:
                    try:
                        batch_index, batch = work_queue.get_nowait()
                    except queue.Empty:
                        break
                    try:
                        if self.stop_requested:
                            break
                        rows, page = self.check_product_batch(context, page, batch, batch_index, PlaywrightError)
                        for row in rows:
                            completed = self.record_result(row, completed_count_safe_index(row, self.products), total, events)
                            events.put(("result", completed, total, row))
                        if not self.stop_requested and completed < total:
                            time.sleep(self.delay)
                    finally:
                        work_queue.task_done()
        except Exception as exc:
            events.put(("error", f"worker {worker_id}: {exc}"))
        finally:
            browser = self.cdp_browsers.pop(worker_id, None)
            if browser is not None:
                try:
                    browser.close()
                except Exception:
                    pass
                process = self.cdp_processes.pop(worker_id, None)
                if process is not None and process.poll() is None:
                    try:
                        process.terminate()
                    except Exception:
                        pass
            elif context is not None:
                try:
                    context.close()
                except Exception:
                    pass
            process = self.cdp_processes.pop(worker_id, None)
            if process is not None and process.poll() is None:
                try:
                    process.terminate()
                except Exception:
                    pass
            if final_status != "bloqueado":
                events.put(("worker_status", worker_id, final_status))

    def run(self, events: queue.Queue) -> None:
        started = time.monotonic()
        try:
            from playwright.sync_api import Error as PlaywrightError
            from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            events.put(("error", f"Playwright no esta instalado: {exc}"))
            return

        self.run_dir.mkdir(parents=True, exist_ok=True)
        events.put(("run_dir", str(self.run_dir)))
        events.put(("phase3", self.worker_count, self.circuit_breaker_threshold))
        for config in self.worker_configs:
            events.put(
                (
                    "worker_config",
                    config.id,
                    config.enabled,
                    config.proxy_configured,
                    config.user_agent_configured,
                )
            )
            if not config.enabled:
                events.put(("worker_status", config.id, "desactivado"))

        work_queue: queue.Queue = queue.Queue()
        batch_size = self.marketplace_config.search_batch_size
        for batch_index, batch in enumerate(chunk_products(self.products, batch_size), start=1):
            work_queue.put((batch_index, batch))

        worker_threads = []
        for config in self.active_worker_configs:
            thread = threading.Thread(
                target=self.browser_worker,
                args=(config.id, work_queue, events, PlaywrightError, PlaywrightTimeoutError, sync_playwright),
                daemon=True,
            )
            worker_threads.append(thread)
            thread.start()

        for thread in worker_threads:
            thread.join()

        elapsed = round(time.monotonic() - started, 1)
        summary = str(self.summary_path)
        if self.circuit_open:
            summary = f"{summary} - corte tecnico activo"
        events.put(("done", elapsed, summary))


class WortenChecker(LeroyChecker):
    marketplace_key = "worten"
    expected_seller_aliases = ("mark jv shop", "newluxgroup", "newlux group")
    seller_search_url = (
        "https://www.worten.pt/search?query=*&facetFilters=seller_id:"
        "e5dae97c-401c-456a-be59-56a4f73b0bb5&utm_source=sellerpage_redirect"
    )

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.block_assets = False
        self.minimal_data_mode = False
        self.data_settle_ms = max(self.data_settle_ms, 2500)
        settings = read_local_settings()
        marketplace_settings = {}
        marketplaces = settings.get("marketplaces", {})
        if isinstance(marketplaces, dict) and isinstance(marketplaces.get("worten"), dict):
            marketplace_settings = marketplaces["worten"]
        self.browser_channel = str(
            marketplace_settings.get("browser_channel")
            or settings.get("worten_browser_channel")
            or "msedge"
        ).strip()
        try:
            configured_timeout = int(
                marketplace_settings.get("challenge_timeout_seconds")
                or settings.get("worten_challenge_timeout_seconds")
                or 600
            )
        except (TypeError, ValueError):
            configured_timeout = 600
        self.challenge_timeout_seconds = max(120, min(configured_timeout, 1800))
        self.manual_cdp_enabled = bool_setting(
            marketplace_settings.get("manual_cdp")
            if "manual_cdp" in marketplace_settings
            else settings.get("worten_manual_cdp"),
            True,
        )
        try:
            self.manual_cdp_port_base = int(
                marketplace_settings.get("manual_cdp_port_base")
                or settings.get("worten_manual_cdp_port_base")
                or 9330
            )
        except (TypeError, ValueError):
            self.manual_cdp_port_base = 9330
        self.manual_cdp_port_base = max(1024, min(self.manual_cdp_port_base, 65000))
        self.cdp_browsers = {}
        self.cdp_processes = {}
        self.pending_challenge_workers: set[int] = set()

    def profile_dir_for_worker(self, worker_id: int) -> Path:
        return FASE_DIR / f"browser_profile_worten_worker_{worker_id}"

    def manual_edge_profile_dir_for_worker(self, worker_id: int) -> Path:
        return FASE_DIR / f"edge_manual_worten_worker_{worker_id}"

    def get_worker_config(self, worker_id: int) -> WorkerConfig:
        for config in self.worker_configs:
            if config.id == worker_id:
                return config
        return WorkerConfig(id=worker_id)

    def find_edge_executable(self) -> str:
        path = shutil.which("msedge.exe") or shutil.which("msedge")
        if path:
            return path
        candidates = [
            Path("C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe"),
            Path("C:/Program Files/Microsoft/Edge/Application/msedge.exe"),
        ]
        for candidate in candidates:
            if candidate.exists():
                return str(candidate)
        raise RuntimeError("No encuentro Microsoft Edge para modo manual CDP")

    def start_worten_manual_edge(self, worker_id: int, port: int, config: WorkerConfig) -> None:
        profile_dir = self.manual_edge_profile_dir_for_worker(worker_id)
        profile_dir.mkdir(parents=True, exist_ok=True)
        args = [
            self.find_edge_executable(),
            f"--remote-debugging-port={port}",
            f"--user-data-dir={profile_dir}",
            "--no-first-run",
            "--no-default-browser-check",
            "--new-window",
        ]
        if config.proxy_server.strip():
            args.append(f"--proxy-server={config.proxy_server.strip()}")
        if config.user_agent.strip():
            args.append(f"--user-agent={config.user_agent.strip()}")
        args.append(self.seller_search_url)
        self.cdp_processes[worker_id] = subprocess.Popen(args, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    def connect_worten_manual_edge(self, p, worker_id: int, config: WorkerConfig):
        port = min(self.manual_cdp_port_base + worker_id - 1, 65535)
        endpoint = f"http://127.0.0.1:{port}"
        try:
            stale_browser = p.chromium.connect_over_cdp(endpoint, timeout=1000)
            stale_browser.close()
            time.sleep(1)
        except Exception:
            pass
        self.start_worten_manual_edge(worker_id, port, config)
        deadline = time.monotonic() + 30
        last_error = None
        while time.monotonic() < deadline:
            try:
                browser = p.chromium.connect_over_cdp(endpoint, timeout=5000)
                if not browser.contexts:
                    raise RuntimeError("Edge CDP sin contexto disponible")
                self.cdp_browsers[worker_id] = browser
                return browser.contexts[0], f"edge-cdp:{port}"
            except Exception as exc:
                last_error = exc
                time.sleep(1)
        raise RuntimeError(f"No puedo conectar con Edge manual en {endpoint}: {last_error}")

    def launch_worten_browser_context(self, p, worker_id: int):
        config = self.get_worker_config(worker_id)
        if self.manual_cdp_enabled:
            return self.connect_worten_manual_edge(p, worker_id, config)

        launch_options = {
            "user_data_dir": str(self.profile_dir_for_worker(worker_id)),
            "headless": False,
            "viewport": {"width": 1365, "height": 900},
            "locale": "pt-PT",
            "timezone_id": "Europe/Lisbon",
            "args": ["--start-minimized"],
        }
        if config.user_agent.strip():
            launch_options["user_agent"] = config.user_agent.strip()
        if config.proxy_server.strip():
            proxy = {"server": config.proxy_server.strip()}
            if config.proxy_username.strip():
                proxy["username"] = config.proxy_username.strip()
            if config.proxy_password.strip():
                proxy["password"] = config.proxy_password.strip()
            launch_options["proxy"] = proxy

        channels = []
        if self.browser_channel:
            channels.append(self.browser_channel)
        channels.extend(["msedge", "chrome", "chromium"])

        last_error = None
        for channel in dict.fromkeys(channels):
            options = dict(launch_options)
            if channel != "chromium":
                options["channel"] = channel
            try:
                return p.chromium.launch_persistent_context(**options), channel
            except Exception as exc:
                last_error = exc
        if last_error:
            raise last_error
        return p.chromium.launch_persistent_context(**launch_options), "chromium"

    def launch_worker_context(self, p, worker_id: int, PlaywrightTimeoutError):
        context, browser_channel = self.launch_worten_browser_context(p, worker_id)
        page = context.pages[0] if context.pages else context.new_page()
        if not self.manual_cdp_enabled:
            minimize_chromium_window(context, page)
        page.set_default_timeout(self.retry_timeout_ms)
        try:
            page.goto(self.seller_search_url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
        except PlaywrightTimeoutError:
            pass
        self.settle_and_stop(page)
        if self.resolve_manual_challenge(context, page, worker_id, "warm-up"):
            minimize_chromium_window(context, page)
            return context, page, True, f"warm-up correcto ({browser_channel})"
        return context, page, False, f"challenge en warm-up Worten ({browser_channel})"

    def settle_and_stop(self, page) -> None:
        try:
            page.wait_for_load_state("load", timeout=min(self.retry_timeout_ms, 12000))
        except Exception:
            pass

    def worten_content_has_listing_signal(self, content: str, eans: list[str] | None = None) -> bool:
        lowered = content[:300000].lower()
        text = visible_text(content[:300000])
        if eans and any(ean in content for ean in eans):
            return True
        markers = (
            "adicionar ao carrinho",
            "juntar ao carrinho",
            "vendido por",
            "mark jv shop",
            "newluxgroup",
            "product-card",
            "resultados",
            "ordenar",
            "sem resultados",
            "nenhum resultado",
            "nao encontramos",
            "n\u00e3o encontramos",
        )
        return any(marker in lowered or marker in text for marker in markers)

    def wait_for_worten_results(self, page, eans: list[str] | None = None, timeout_seconds: int = 18) -> tuple[str, str, str]:
        deadline = time.monotonic() + max(3, timeout_seconds)
        last_content = ""
        last_title = ""
        last_url = ""
        challenge_hits = 0
        while time.monotonic() < deadline and not self.stop_requested:
            try:
                last_content = page.content()
                last_title = page.title()
                last_url = page.url
            except Exception:
                time.sleep(0.5)
                continue
            if self.worten_content_has_listing_signal(last_content, eans):
                return last_content, last_title, last_url
            if self.worten_challenge_detected(last_content, last_title, last_url):
                challenge_hits += 1
                if challenge_hits >= 3:
                    return last_content, last_title, last_url
            else:
                challenge_hits = 0
            try:
                page.wait_for_timeout(1000)
            except Exception:
                time.sleep(1)
        return last_content, last_title, last_url

    def collect_worten_paginated_results(self, page, eans: list[str], PlaywrightError, max_pages: int = 21) -> tuple[str, str, str]:
        contents: list[str] = []
        content, title, final_url = self.wait_for_worten_results(page, eans, timeout_seconds=18)
        if content:
            contents.append(content)
        found = {ean for ean in eans if ean in content}
        if self.worten_challenge_detected(content, title, final_url) or len(found) >= len(eans):
            return "\n".join(contents), title, final_url

        base_url = final_url
        for page_number in range(2, max_pages + 1):
            if len(found) >= len(eans) or self.stop_requested:
                break
            page_url = re.sub(r"([?&])page=\d+", "", base_url).rstrip("?&")
            separator = "&" if "?" in page_url else "?"
            page_url = f"{page_url}{separator}page={page_number}"
            try:
                page.goto(page_url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
            except PlaywrightError:
                break
            page_content, page_title, page_final_url = self.wait_for_worten_results(page, eans, timeout_seconds=15)
            if self.worten_challenge_detected(page_content, page_title, page_final_url):
                combined = "\n".join(contents)
                if combined and self.worten_content_has_listing_signal(combined, eans):
                    return combined, title, final_url
                return page_content, page_title, page_final_url
            if not self.worten_content_has_listing_signal(page_content, eans):
                break
            contents.append(page_content)
            found.update(ean for ean in eans if ean in page_content)
            title = page_title or title
            final_url = page_final_url or final_url
        return "\n".join(contents), title, final_url

    def worten_challenge_detected(self, content: str, title: str, url: str) -> bool:
        lowered = f"{title}\n{url}\n{content[:250000]}".lower()
        text = visible_text(content[:250000])
        strong_result_signal = (
            "vendido por" in text
            and ("mark jv shop" in text or "newluxgroup" in text or "newlux group" in text)
            and ("€" in content or "&euro;" in lowered or re.search(r"\b[0-9]+[,.][0-9]{2}\b", text))
        )
        if strong_result_signal:
            return False
        has_usable_worten_signal = "worten" in lowered and (
            "adicionar ao carrinho" in text
            or "juntar ao carrinho" in text
            or "vendido por" in text
            or "mark jv shop" in text
            or "newluxgroup" in text
            or "pre\u00e7o" in text
            or "preco" in text
            or "ordenar" in text
            or "resultados" in text
            or "product-card" in lowered
            or "seller_id" in lowered
        )
        visible_challenge_markers = [
            "executando verifica\u00e7\u00e3o de seguran\u00e7a",
            "executando verificacao de seguranca",
            "confirme que \u00e9 humano",
            "confirme que e humano",
            "esta p\u00e1gina \u00e9 exibida enquanto",
            "esta pagina e exibida enquanto",
            "servi\u00e7o de seguran\u00e7a para prote\u00e7\u00e3o contra bots",
            "servico de seguranca para protecao contra bots",
            "checking your browser",
            "verify you are human",
            "verifying you are human",
            "human verification",
            "um momento",
            "access denied",
            "forbidden",
            "blocked",
            "bloqueado",
        ]
        visible_challenge_in_text = any(marker in text for marker in visible_challenge_markers)
        visible_challenge_in_title = any(marker in title.lower() for marker in visible_challenge_markers)
        if has_usable_worten_signal and not visible_challenge_in_text:
            return False
        if visible_challenge_in_text or (visible_challenge_in_title and not has_usable_worten_signal):
            return True
        return False

    def is_circuit_signal(self, row: DashboardResult) -> bool:
        evidence = f"{row.status} {row.reason} {row.title} {row.final_url}".casefold()
        return row.challenge_detected or "challenge" in evidence or "captcha" in evidence

    def ensure_worten_ready_after_challenge(self, page) -> None:
        try:
            content = page.content()
            title = page.title()
            final_url = page.url
        except Exception:
            return
        if self.worten_challenge_detected(content, title, final_url):
            return
        if "search" not in final_url.lower():
            try:
                page.goto(self.seller_search_url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
            except Exception:
                pass
        try:
            page.wait_for_load_state("load", timeout=min(self.retry_timeout_ms, 12000))
        except Exception:
            pass
        try:
            page.wait_for_timeout(self.data_settle_ms)
        except Exception:
            pass

    def resolve_manual_challenge(self, context, page, worker_id: int, reason: str, timeout_seconds: int | None = None) -> bool:
        timeout_seconds = timeout_seconds or self.challenge_timeout_seconds
        try:
            content = page.content()
            title = page.title()
            final_url = page.url
        except Exception:
            return False
        if not self.worten_challenge_detected(content, title, final_url):
            if worker_id in self.pending_challenge_workers:
                self.pending_challenge_workers.discard(worker_id)
                if self.events:
                    self.events.put(("human_challenge_resolved", worker_id))
                    self.events.put(("worker_status", worker_id, "activo"))
            return True

        if self.events:
            self.pending_challenge_workers.add(worker_id)
            self.events.put(("worker_status", worker_id, "pendiente challenge"))
            self.events.put(("human_challenge_required", worker_id, reason, final_url))
        show_chromium_window(context, page)
        deadline = time.monotonic() + timeout_seconds
        while time.monotonic() < deadline and not self.stop_requested:
            try:
                page.wait_for_timeout(3000)
                content = page.content()
                title = page.title()
                final_url = page.url
            except Exception:
                continue
            if not self.worten_challenge_detected(content, title, final_url):
                self.ensure_worten_ready_after_challenge(page)
                minimize_chromium_window(context, page)
                self.pending_challenge_workers.discard(worker_id)
                if self.events:
                    self.events.put(("human_challenge_resolved", worker_id))
                    self.events.put(("worker_status", worker_id, "activo"))
                return True
        if self.events:
            self.events.put(("human_challenge_timeout", worker_id))
        return False

    def seller_matches(self, seller: str) -> bool:
        seller_norm = normalize(seller)
        seller_compact = re.sub(r"[^a-z0-9]+", "", seller_norm)
        expected_values = [self.expected_seller] if self.expected_seller.strip() else []
        expected_values.extend(self.expected_seller_aliases)
        for expected in expected_values:
            expected_norm = normalize(expected)
            expected_compact = re.sub(r"[^a-z0-9]+", "", expected_norm)
            if seller_norm == expected_norm or seller_compact == expected_compact:
                return True
        return False

    def extract_search_card(self, content: str, product: InputProduct) -> WortenSearchCard | None:
        ean = re.escape(product.ean)
        patterns = [
            rf"<article\b[\s\S]*?{ean}[\s\S]*?</article>",
            rf"<li\b[\s\S]*?{ean}[\s\S]*?</li>",
            rf"<div\b[\s\S]*?{ean}[\s\S]*?</div>",
        ]
        block = ""
        for pattern in patterns:
            match = re.search(pattern, content, flags=re.IGNORECASE)
            if match:
                block = match.group(0)
                break
        if not block and product.ean in content:
            index = content.find(product.ean)
            block = content[max(0, index - 12000) : index + 12000]
        if not block:
            return None

        url = extract_first(r'href=["\'](https://www\.worten\.pt/[^"\']+)["\']', block)
        if not url:
            relative = extract_first(r'href=["\'](/[^"\']+/p/[^"\']+)["\']', block)
            if relative:
                url = "https://www.worten.pt" + relative

        title_text = extract_first(r'<img[^>]+alt=["\']([^"\']+)["\']', block)
        if not title_text:
            title_text = extract_first(r'<a[^>]+title=["\']([^"\']+)["\']', block)
        if not title_text:
            title_text = visible_text(block)[:180]

        price = extract_first(r'([0-9]+(?:[,.][0-9]{2})\s*(?:\u20ac|&euro;|EUR))', block)
        seller = extract_first(r"(?:Vendido por|Seller|Vendedor)\s*<[^>]*>\s*([^<]+)", block)
        if not seller:
            for alias in self.expected_seller_aliases:
                if alias in normalize(block):
                    seller = alias
                    break
        add_to_cart = any(marker in normalize(block) for marker in ("adicionar", "juntar ao carrinho", "comprar", "adicionar ao carrinho"))
        return WortenSearchCard(product.ean, url, title_text, price, seller, add_to_cart)

    def submit_search_input(self, page, eans: list[str]) -> bool:
        query = " ".join(eans)
        selectors = [
            "input[type='search']",
            "input[name='query']",
            "input[name='q']",
            "input[placeholder*='Procurar']",
            "input[placeholder*='Pesquisar']",
            "input[placeholder*='Search']",
            "[data-testid*='search'] input",
            "#search-input",
        ]
        for selector in selectors:
            try:
                locator = page.locator(selector).first
                locator.wait_for(state="visible", timeout=15000)
                locator.click(timeout=5000)
                try:
                    locator.fill("")
                except Exception:
                    page.keyboard.press("Control+A")
                    page.keyboard.press("Backspace")
                locator.fill(query)
                page.keyboard.press("Enter")
                return True
            except Exception:
                continue
        return False

    def search_batch_via_input(self, context, page, batch: list[InputProduct], worker_id: int, PlaywrightError) -> tuple[str, str, str, object]:
        eans = [product.ean for product in batch]
        try:
            page.goto(self.seller_search_url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
        except PlaywrightError:
            pass
        self.settle_and_stop(page)
        if not self.resolve_manual_challenge(context, page, worker_id, "busqueda Worten"):
            raise ManualChallengeNotResolved("challenge Worten pendiente/no resuelto antes de buscar")
        submitted = self.submit_search_input(page, eans)
        if not submitted:
            return page.content(), "Worten search input not found", page.url, page
        self.settle_and_stop(page)
        content, title, final_url = self.collect_worten_paginated_results(page, eans, PlaywrightError, max_pages=21)
        return content, title, final_url, page

    def classify_batch(
        self,
        batch: list[InputProduct],
        content: str,
        title: str,
        final_url: str,
        elapsed_seconds: float,
        html_path: str,
    ) -> list[DashboardResult]:
        is_challenge = self.worten_challenge_detected(content, title, final_url)
        rows: list[DashboardResult] = []
        for product in batch:
            card = self.extract_search_card(content, product)
            if is_challenge:
                light = "gray"
                status = "INCIERTA"
                reason = "challenge detectado"
                seller = ""
                price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"
            elif len(content.encode("utf-8")) < 10_000:
                light = "gray"
                status = "INCIERTA"
                reason = "HTML demasiado pequeno"
                seller = ""
                price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"
            elif card:
                seller = card.seller
                price = card.price
                result_url = card.url or final_url
                product_title = card.title or title
                offer_available = card.add_to_cart or bool(card.price)
                availability = "seller_search_result"
                if seller:
                    buybox_ok = "yes" if self.seller_matches(seller) else "no"
                else:
                    buybox_ok = "yes"
                    seller = "Mark JV shop"
                if buybox_ok == "no":
                    light = "yellow"
                    status = "BUYBOX_PERDIDA"
                    reason = f"seller actual={seller}"
                elif not offer_available:
                    light = "yellow"
                    status = "NO_DISPONIBLE_CON_STOCK"
                    reason = "producto Worten encontrado sin senal de compra/precio"
                else:
                    light = "green"
                    status = "OK"
                    reason = "oferta Worten encontrada en pagina dedicada del seller"
            else:
                text = visible_text(content)
                if any(marker in text for marker in ("sem resultados", "no results", "nenhum resultado", "nao encontramos", "n\u00e3o encontramos")):
                    light = "red"
                    status = "NO_VIVA"
                    reason = "EAN no encontrado en pagina dedicada del seller Worten"
                else:
                    light = "gray"
                    status = "INCIERTA"
                    reason = "Worten sin senal clara de producto ni de no resultados"
                seller = ""
                price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"

            rows.append(
                DashboardResult(
                    ean=product.ean,
                    reference=product.reference,
                    feed_quantity=product.quantity,
                    feed_price=product.price,
                    light=light,
                    status=status,
                    final_url=result_url,
                    title=product_title,
                    sku=product.ean if card else "",
                    seller_name=seller,
                    marketplace_price=price,
                    product_availability=availability,
                    offer_available=offer_available,
                    buybox_ok=buybox_ok,
                    challenge_detected=is_challenge,
                    reason=reason,
                    elapsed_seconds=elapsed_seconds,
                    html_path=html_path if light != "green" else "",
                )
            )
        return rows

    def check_product_batch(
        self,
        context,
        page,
        batch: list[InputProduct],
        batch_index: int,
        worker_id: int,
        PlaywrightError,
    ) -> tuple[list[DashboardResult], object]:
        item_started = time.monotonic()
        try:
            content, title, final_url, page = self.search_batch_via_input(context, page, batch, worker_id, PlaywrightError)
        except PlaywrightError as exc:
            elapsed = round(time.monotonic() - item_started, 2)
            rows = []
            for product in batch:
                rows.append(
                    DashboardResult(
                        ean=product.ean,
                        reference=product.reference,
                        feed_quantity=product.quantity,
                        feed_price=product.price,
                        light="gray",
                        status="INCIERTA",
                        final_url=self.seller_search_url,
                        title="",
                        sku="",
                        seller_name="",
                        marketplace_price="",
                        product_availability="unknown",
                        offer_available=False,
                        buybox_ok="unknown",
                        challenge_detected=False,
                        reason=f"error de navegacion Worten batch: {exc}",
                        elapsed_seconds=elapsed,
                        html_path="",
                    )
                )
            return rows, page
        has_listing = self.worten_content_has_listing_signal(content or "", [product.ean for product in batch])
        if self.worten_challenge_detected(content or "", title or "", final_url or self.seller_search_url) and not has_listing:
            if not self.resolve_manual_challenge(context, page, worker_id, "resultado Worten"):
                elapsed = round(time.monotonic() - item_started, 2)
                rows = []
                for product in batch:
                    rows.append(
                        DashboardResult(
                            ean=product.ean,
                            reference=product.reference,
                            feed_quantity=product.quantity,
                            feed_price=product.price,
                            light="gray",
                            status="INCIERTA",
                            final_url=final_url or self.seller_search_url,
                            title=title or "",
                            sku="",
                            seller_name="",
                            marketplace_price="",
                            product_availability="unknown",
                            offer_available=False,
                            buybox_ok="unknown",
                            challenge_detected=True,
                            reason="challenge visible sin resultados tras busqueda Worten",
                            elapsed_seconds=elapsed,
                            html_path="",
                        )
                    )
                return rows, page
            content, title, final_url = page.content(), page.title(), page.url
        elapsed = round(time.monotonic() - item_started, 2)
        html_file = self.run_dir / f"worten_batch_{batch_index:04d}.html"
        if content:
            html_file.write_text(content, encoding="utf-8")
        rows = self.classify_batch(batch, content or "", title or "", final_url or self.seller_search_url, elapsed, str(html_file))
        return rows, page

    def browser_worker(
        self,
        worker_id: int,
        work_queue: queue.Queue,
        events: queue.Queue,
        PlaywrightError,
        PlaywrightTimeoutError,
        sync_playwright,
    ) -> None:
        context = None
        final_status = "cerrado"
        try:
            with sync_playwright() as p:
                events.put(("worker_status", worker_id, "calentando"))
                context, page, healthy, warmup_reason = self.launch_worker_context(p, worker_id, PlaywrightTimeoutError)
                if not healthy:
                    final_status = "bloqueado"
                    events.put(("worker_status", worker_id, "bloqueado"))
                    events.put(("worker_blocked", worker_id, warmup_reason))
                    self.notify_worker_blocked(worker_id, warmup_reason, page, events)
                    return
                events.put(("worker_status", worker_id, "activo"))
                total = len(self.products)
                while not self.stop_requested:
                    try:
                        batch_index, batch = work_queue.get_nowait()
                    except queue.Empty:
                        break
                    try:
                        if self.stop_requested:
                            break
                        try:
                            rows, page = self.check_product_batch(context, page, batch, batch_index, worker_id, PlaywrightError)
                        except ManualChallengeNotResolved as exc:
                            final_status = "bloqueado"
                            events.put(("worker_status", worker_id, "bloqueado"))
                            events.put(("worker_blocked", worker_id, str(exc)))
                            self.notify_worker_blocked(worker_id, str(exc), page, events)
                            return
                        events.put(("worker_status", worker_id, "activo"))
                        for row in rows:
                            completed = self.record_result(row, completed_count_safe_index(row, self.products), total, events)
                            events.put(("result", completed, total, row))
                        if not self.stop_requested and completed < total:
                            time.sleep(self.delay)
                    finally:
                        work_queue.task_done()
        except Exception as exc:
            self.stop_requested = True
            events.put(("error", f"worker {worker_id}: {exc}"))
        finally:
            browser = self.cdp_browsers.pop(worker_id, None)
            if browser is not None:
                try:
                    browser.close()
                except Exception:
                    pass
                process = self.cdp_processes.pop(worker_id, None)
                if process is not None and process.poll() is None:
                    try:
                        process.terminate()
                    except Exception:
                        pass
            elif context is not None:
                try:
                    context.close()
                except Exception:
                    pass
            process = self.cdp_processes.pop(worker_id, None)
            if process is not None and process.poll() is None:
                try:
                    process.terminate()
                except Exception:
                    pass
            if final_status != "bloqueado":
                events.put(("worker_status", worker_id, final_status))

    def run(self, events: queue.Queue) -> None:
        self.events = events
        started = time.monotonic()
        try:
            from playwright.sync_api import Error as PlaywrightError
            from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            events.put(("error", f"Playwright no esta instalado: {exc}"))
            return

        self.run_dir.mkdir(parents=True, exist_ok=True)
        events.put(("run_dir", str(self.run_dir)))
        events.put(("phase3", self.worker_count, self.circuit_breaker_threshold))
        for config in self.worker_configs:
            events.put(
                (
                    "worker_config",
                    config.id,
                    config.enabled,
                    config.proxy_configured,
                    config.user_agent_configured,
                )
            )
            if not config.enabled:
                events.put(("worker_status", config.id, "desactivado"))

        work_queue: queue.Queue = queue.Queue()
        batch_size = self.marketplace_config.search_batch_size
        for batch_index, batch in enumerate(chunk_products(self.products, batch_size), start=1):
            work_queue.put((batch_index, batch))

        worker_threads = []
        for config in self.active_worker_configs:
            thread = threading.Thread(
                target=self.browser_worker,
                args=(config.id, work_queue, events, PlaywrightError, PlaywrightTimeoutError, sync_playwright),
                daemon=True,
            )
            worker_threads.append(thread)
            thread.start()

        for thread in worker_threads:
            thread.join()

        elapsed = round(time.monotonic() - started, 1)
        summary = str(self.summary_path)
        if self.circuit_open:
            summary = f"{summary} - corte tecnico activo"
        events.put(("done", elapsed, summary))


WortenCheckerLegacy = WortenChecker


class WortenChecker(WortenCheckerLegacy):
    marketplace_key = "worten"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.data_settle_ms = max(self.data_settle_ms, 5000)
        self.worten_cookie_checked_workers: set[int] = set()

    def profile_dir_for_worker(self, worker_id: int) -> Path:
        return FASE_DIR / f"browser_profile_worten_clean_worker_{worker_id}"

    def manual_edge_profile_dir_for_worker(self, worker_id: int) -> Path:
        return FASE_DIR / f"edge_manual_worten_clean_worker_{worker_id}"

    def start_worten_manual_edge(self, worker_id: int, port: int, config: WorkerConfig) -> None:
        profile_dir = self.manual_edge_profile_dir_for_worker(worker_id)
        profile_dir.mkdir(parents=True, exist_ok=True)
        args = [
            self.find_edge_executable(),
            f"--remote-debugging-port={port}",
            f"--user-data-dir={profile_dir}",
            "--disable-extensions",
            "--disable-component-extensions-with-background-pages",
            "--disable-background-networking",
            "--disable-sync",
            "--disable-features=msEdgeAccountExtension,msSingleSignOnOSForPrimaryAccountIsShared,EdgeSignIn,EnableSyncConsent",
            "--no-first-run",
            "--no-default-browser-check",
            "--guest",
            "--new-window",
        ]
        if config.proxy_server.strip():
            args.append(f"--proxy-server={config.proxy_server.strip()}")
        if config.user_agent.strip():
            args.append(f"--user-agent={config.user_agent.strip()}")
        args.append(self.seller_search_url)
        self.cdp_processes[worker_id] = subprocess.Popen(args, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    def settle_and_stop(self, page) -> None:
        try:
            page.wait_for_load_state("load", timeout=min(self.retry_timeout_ms, 15000))
        except Exception:
            pass
        try:
            page.wait_for_timeout(self.data_settle_ms)
        except Exception:
            pass

    def accept_worten_cookies(self, page, attempts: int = 6) -> bool:
        text_pattern = re.compile(
            r"(aceitar\s+(todos|todas|tudo|cookies|e\s+fechar|e\s+continuar)|"
            r"aceito|concordo|permitir\s+todos|accept\s+all|allow\s+all|i\s+agree)",
            re.IGNORECASE,
        )
        selectors = (
            "#onetrust-accept-btn-handler",
            "#accept-recommended-btn-handler",
            "#didomi-notice-agree-button",
            "#CybotCookiebotDialogBodyLevelButtonLevelOptinAllowAll",
            "#truste-consent-button",
            "button[data-testid*='accept']",
            "button[id*='accept' i]",
            "button[id*='agree' i]",
            "button[class*='accept' i]",
            "[role='button'][id*='accept' i]",
        )
        for _attempt in range(max(1, attempts)):
            frames = [page]
            try:
                frames.extend(page.frames)
            except Exception:
                pass
            for frame in frames:
                for selector in selectors:
                    try:
                        locator = frame.locator(selector).first
                        if locator.is_visible(timeout=700):
                            locator.click(timeout=2000, force=True)
                            page.wait_for_timeout(900)
                            return True
                    except Exception:
                        continue
                try:
                    button = frame.get_by_role("button", name=text_pattern).first
                    if button.is_visible(timeout=700):
                        button.click(timeout=2000, force=True)
                        page.wait_for_timeout(900)
                        return True
                except Exception:
                    pass
                try:
                    clicked = frame.evaluate(
                        """
                        () => {
                          const words = [
                            'aceitar tudo', 'aceitar todos', 'aceitar todas', 'aceitar cookies',
                            'aceitar e fechar', 'aceitar e continuar', 'aceito', 'concordo',
                            'permitir todos', 'accept all', 'allow all', 'i agree'
                          ];
                          const scopeRegex = /cookie|cookies|consent|privacidade|preferencia|preferência|onetrust|didomi|cookiebot/i;
                          const textOf = (node) => (node && (node.innerText || node.textContent) || '').trim().toLowerCase();
                          const candidates = Array.from(document.querySelectorAll('button, [role="button"], input[type="button"], input[type="submit"], a'));
                          for (const el of candidates) {
                            const text = textOf(el) || String(el.value || '').trim().toLowerCase();
                            if (!words.some(word => text.includes(word))) continue;
                            const scope = el.closest('[id*="cookie" i], [class*="cookie" i], [id*="consent" i], [class*="consent" i], [id*="onetrust" i], [class*="onetrust" i], [id*="didomi" i], [class*="didomi" i], [id*="Cookiebot" i]') || el.parentElement;
                            const scopeText = textOf(scope);
                            if (scopeRegex.test(scopeText) || scopeRegex.test(String(el.id || '')) || scopeRegex.test(String(el.className || ''))) {
                              el.click();
                              return true;
                            }
                          }
                          return false;
                        }
                        """
                    )
                    if clicked:
                        page.wait_for_timeout(900)
                        return True
                except Exception:
                    pass
            try:
                page.wait_for_timeout(1000)
            except Exception:
                time.sleep(1)
        button_names = (
            r"aceitar todos",
            r"aceitar todas",
            r"aceitar tudo",
            r"aceitar cookies",
            r"aceitar",
            r"permitir todos",
            r"concordo",
        )
        for name in button_names:
            try:
                button = page.get_by_role("button", name=re.compile(name, re.IGNORECASE)).first
                if button.is_visible(timeout=1000):
                    button.click(timeout=2000)
                    page.wait_for_timeout(600)
                    return True
            except Exception:
                continue
        try:
            page.evaluate(
                """
                () => {
                  const words = ['aceitar todos', 'aceitar todas', 'aceitar cookies', 'permitir todos', 'concordo'];
                  const buttons = Array.from(document.querySelectorAll('button, [role="button"]'));
                  for (const button of buttons) {
                    const text = (button.innerText || button.textContent || '').trim().toLowerCase();
                    const scope = (button.closest('[id*="cookie"], [class*="cookie"], [id*="consent"], [class*="consent"], [id*="onetrust"], [class*="onetrust"]') || button.parentElement);
                    const scopeText = ((scope && (scope.innerText || scope.textContent)) || '').toLowerCase();
                    if (words.some(word => text.includes(word)) && /cookie|consent|privacidade|preferência|preferencia/.test(scopeText)) {
                      button.click();
                      return true;
                    }
                  }
                  return false;
                }
                """
            )
            page.wait_for_timeout(600)
        except Exception:
            pass
        return False

    def ensure_worten_cookies(self, page, worker_id: int, attempts: int = 1, force: bool = False) -> bool:
        if not force and worker_id in self.worten_cookie_checked_workers:
            return False
        clicked = self.accept_worten_cookies(page, attempts=attempts)
        if clicked or force:
            self.worten_cookie_checked_workers.add(worker_id)
        return clicked

    def resolve_manual_challenge(self, context, page, worker_id: int, reason: str, timeout_seconds: int | None = None) -> bool:
        had_challenge = False
        try:
            had_challenge = self.worten_challenge_detected(page.content(), page.title(), page.url)
        except Exception:
            pass
        resolved = super().resolve_manual_challenge(context, page, worker_id, reason, timeout_seconds)
        if resolved:
            self.ensure_worten_cookies(page, worker_id, attempts=10 if had_challenge else 1, force=had_challenge)
        return resolved

    def worten_challenge_detected(self, content: str, title: str, url: str) -> bool:
        text = visible_text(content[:250000])
        lowered = f"{title}\n{url}\n{text}".casefold()
        result_markers = (
            "vendido por",
            "mark jv shop",
            "resultados de pesquisa",
            "ordenar por",
            "1000+ produtos",
            "produtos",
            "seller_id",
        )
        challenge_markers = (
            "executando verifica\u00e7\u00e3o de seguran\u00e7a",
            "executando verificacao de seguranca",
            "confirme que \u00e9 humano",
            "confirme que e humano",
            "esta pagina e exibida enquanto",
            "esta p\u00e1gina \u00e9 exibida enquanto",
            "servico de seguranca para protecao contra bots",
            "servi\u00e7o de seguran\u00e7a para prote\u00e7\u00e3o contra bots",
            "checking your browser",
            "verify you are human",
            "verifying you are human",
            "human verification",
            "um momento",
        )
        if any(marker in lowered for marker in result_markers):
            return False
        return any(marker in lowered for marker in challenge_markers)

    def submit_search_input(self, page, eans: list[str]) -> bool:
        query = " ".join(eans[: self.marketplace_config.search_batch_size])
        selectors = [
            "input[type='search']",
            "input[name='query']",
            "input[name='q']",
            "input[placeholder*='Pesquisa']",
            "input[placeholder*='Procurar']",
            "input[placeholder*='Pesquisar']",
            "[data-testid*='search'] input",
            "#search-input",
        ]
        for selector in selectors:
            try:
                locator = page.locator(selector).first
                locator.wait_for(state="visible", timeout=15000)
                locator.click(timeout=5000)
                try:
                    locator.fill("")
                except Exception:
                    page.keyboard.press("Control+A")
                    page.keyboard.press("Backspace")
                locator.fill(query)
                page.keyboard.press("Enter")
                return True
            except Exception:
                continue
        return False

    def ensure_search_input_available(self, context, page, worker_id: int, PlaywrightError) -> bool:
        self.ensure_worten_cookies(page, worker_id, attempts=1)
        if self.resolve_manual_challenge(context, page, worker_id, "busqueda Worten"):
            self.ensure_worten_cookies(page, worker_id, attempts=1)
        else:
            return False
        selectors = [
            "input[type='search']",
            "input[name='query']",
            "input[name='q']",
            "input[placeholder*='Pesquisa']",
            "input[placeholder*='Procurar']",
            "input[placeholder*='Pesquisar']",
            "[data-testid*='search'] input",
            "#search-input",
        ]
        for selector in selectors:
            try:
                page.locator(selector).first.wait_for(state="visible", timeout=2500)
                return True
            except Exception:
                continue

        try:
            page.goto(self.seller_search_url, wait_until="domcontentloaded", timeout=self.nav_timeout_ms)
        except PlaywrightError:
            pass
        self.settle_and_stop(page)
        self.ensure_worten_cookies(page, worker_id, attempts=3, force=True)
        if not self.resolve_manual_challenge(context, page, worker_id, "recuperar buscador Worten"):
            return False
        self.ensure_worten_cookies(page, worker_id, attempts=3, force=True)
        for selector in selectors:
            try:
                page.locator(selector).first.wait_for(state="visible", timeout=5000)
                return True
            except Exception:
                continue
        return False

    def wait_for_worten_results(self, page, eans: list[str] | None = None, timeout_seconds: int = 24) -> tuple[str, str, str]:
        deadline = time.monotonic() + max(8, timeout_seconds)
        last_content = ""
        last_title = ""
        last_url = ""
        requested = eans or []
        while time.monotonic() < deadline and not self.stop_requested:
            try:
                last_content = page.content()
                last_title = page.title()
                last_url = page.url
                text = page.locator("body").inner_text(timeout=4000)
            except Exception:
                time.sleep(0.75)
                continue
            lowered_text = text.casefold()
            if self.worten_challenge_detected(last_content, last_title, last_url):
                return last_content, last_title, last_url
            query_has_requested = not requested or all(ean in last_url for ean in requested)
            has_result_shell = "resultados de pesquisa" in lowered_text or "ordenar por" in lowered_text or "produtos" in lowered_text
            has_seller = "vendido por" in lowered_text or "mark jv" in lowered_text
            has_requested_product_url = any(f"mrkean-{ean}" in last_content for ean in requested)
            if has_requested_product_url or (has_result_shell and has_seller):
                if query_has_requested:
                    return last_content, last_title, last_url
            try:
                page.wait_for_timeout(800)
            except Exception:
                time.sleep(0.8)
        return last_content, last_title, last_url

    def search_batch_via_input(self, context, page, batch: list[InputProduct], worker_id: int, PlaywrightError) -> tuple[str, str, str, object]:
        eans = [product.ean for product in batch[: self.marketplace_config.search_batch_size]]
        if not self.ensure_search_input_available(context, page, worker_id, PlaywrightError):
            raise ManualChallengeNotResolved("challenge Worten pendiente/no resuelto antes de buscar")
        if not self.submit_search_input(page, eans):
            if not self.ensure_search_input_available(context, page, worker_id, PlaywrightError):
                return page.content(), "Worten search input not found", page.url, page
            if not self.submit_search_input(page, eans):
                return page.content(), "Worten search input not found", page.url, page
        content, title, final_url = self.wait_for_worten_results(page, eans, timeout_seconds=24)
        try:
            page.wait_for_timeout(2500)
        except Exception:
            pass
        return content, title, final_url, page

    def parse_worten_card(self, raw: dict) -> WortenSearchCard:
        text = str(raw.get("text", "") or "")
        compact_text = re.sub(r"\s+", " ", text)
        url = str(raw.get("url", "") or "")
        ean = extract_first(r"mrkean-(\d{8,14})", url)
        if not ean:
            ean = extract_first(r"\b(\d{8,14})\b", text)
        seller = extract_first(r"Vendido por\s+([^\n]+)", text)
        price = ""
        price_match = re.search(r"\u20ac\s*([0-9]+)\s*,\s*([0-9]{2})", compact_text)
        if price_match:
            price = f"{price_match.group(1)},{price_match.group(2)}"
        else:
            price = extract_first(r"([0-9]+[,.][0-9]{2})\s*\u20ac", compact_text)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        title = str(raw.get("title", "") or "").strip()
        if not title:
            title = next(
                (
                    line
                    for line in lines
                    if len(line) > 8
                    and "\u20ac" not in line
                    and not line.lower().startswith(("iva ", "vendido por", "entrega ", "+", "comparar"))
                ),
                "",
            )
        add_to_cart = any(marker in normalize(text) for marker in ("adicionar", "juntar ao carrinho", "comprar", "cesto"))
        return WortenSearchCard(ean=ean, url=url, title=title, price=price, seller=seller, add_to_cart=add_to_cart)

    def extract_visible_search_cards(self, page) -> list[WortenSearchCard]:
        raw_cards = page.evaluate(
            """
            () => {
              const anchors = Array.from(document.querySelectorAll('a[href*="/produtos/"]'));
              const cards = [];
              for (const anchor of anchors) {
                let node = anchor;
                for (let depth = 0; node && depth < 8; depth += 1) {
                  const text = (node.innerText || '').trim();
                  if (/Vendido por/i.test(text) && /\\u20ac|€/i.test(text)) {
                    const img = node.querySelector('img[alt]');
                    cards.push({
                      url: new URL(anchor.getAttribute('href'), location.href).href,
                      title: img ? img.getAttribute('alt') : '',
                      text
                    });
                    break;
                  }
                  node = node.parentElement;
                }
              }
              return cards;
            }
            """
        )
        cards: list[WortenSearchCard] = []
        for raw in raw_cards:
            if isinstance(raw, dict):
                card = self.parse_worten_card(raw)
                if card.url or card.ean or card.seller or card.price:
                    cards.append(card)
        return cards

    def classify_batch_from_cards(
        self,
        batch: list[InputProduct],
        cards: list[WortenSearchCard],
        content: str,
        title: str,
        final_url: str,
        elapsed_seconds: float,
        html_path: str,
    ) -> list[DashboardResult]:
        is_challenge = self.worten_challenge_detected(content, title, final_url)
        cards_by_ean: dict[str, WortenSearchCard] = {}
        for card in cards:
            if card.ean and card.ean not in cards_by_ean:
                cards_by_ean[card.ean] = card
        text = visible_text(content)
        no_results_signal = any(
            marker in text
            for marker in ("sem resultados", "nenhum resultado", "nao encontramos", "n\u00e3o encontramos")
        )
        rows: list[DashboardResult] = []
        for product in batch:
            card = cards_by_ean.get(product.ean)
            if is_challenge:
                light = "gray"
                status = "INCIERTA"
                reason = "challenge detectado"
                seller = ""
                price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"
            elif card:
                seller = card.seller
                price = card.price
                result_url = card.url or final_url
                product_title = card.title or title
                offer_available = bool(price or seller or result_url)
                availability = "search_result_price_seller"
                buybox_ok = "yes" if self.seller_matches(seller or "MARK JV SHOP") else "no"
                if buybox_ok == "no":
                    light = "yellow"
                    status = "BUYBOX_PERDIDA"
                    reason = f"seller actual={seller}"
                elif not price:
                    light = "yellow"
                    status = "NO_DISPONIBLE_CON_STOCK"
                    reason = "tarjeta Worten encontrada sin precio visible"
                else:
                    light = "green"
                    status = "OK"
                    reason = "oferta Worten encontrada desde busqueda de 10 EAN"
            elif no_results_signal:
                light = "red"
                status = "NO_VIVA"
                reason = "EAN no encontrado en busqueda Worten"
                seller = ""
                price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "no_results"
            else:
                light = "gray"
                status = "INCIERTA"
                reason = "Worten devolvio resultados, pero no se pudo mapear tarjeta al EAN"
                seller = ""
                price = ""
                result_url = final_url
                product_title = title
                offer_available = False
                buybox_ok = "unknown"
                availability = "unknown"

            rows.append(
                DashboardResult(
                    ean=product.ean,
                    reference=product.reference,
                    feed_quantity=product.quantity,
                    feed_price=product.price,
                    light=light,
                    status=status,
                    final_url=result_url,
                    title=product_title,
                    sku=card.ean if card else "",
                    seller_name=seller,
                    marketplace_price=price,
                    product_availability=availability,
                    offer_available=offer_available,
                    buybox_ok=buybox_ok,
                    challenge_detected=is_challenge,
                    reason=reason,
                    elapsed_seconds=elapsed_seconds,
                    html_path=html_path if light != "green" else "",
                )
            )
        return rows

    def check_product_batch(
        self,
        context,
        page,
        batch: list[InputProduct],
        batch_index: int,
        worker_id: int,
        PlaywrightError,
    ) -> tuple[list[DashboardResult], object]:
        item_started = time.monotonic()
        try:
            content, title, final_url, page = self.search_batch_via_input(context, page, batch, worker_id, PlaywrightError)
            cards = self.extract_visible_search_cards(page)
            content = page.content()
            title = page.title()
            final_url = page.url
        except PlaywrightError as exc:
            elapsed = round(time.monotonic() - item_started, 2)
            rows = [
                DashboardResult(
                    ean=product.ean,
                    reference=product.reference,
                    feed_quantity=product.quantity,
                    feed_price=product.price,
                    light="gray",
                    status="INCIERTA",
                    final_url=self.seller_search_url,
                    title="",
                    sku="",
                    seller_name="",
                    marketplace_price="",
                    product_availability="unknown",
                    offer_available=False,
                    buybox_ok="unknown",
                    challenge_detected=False,
                    reason=f"error de navegacion Worten batch: {exc}",
                    elapsed_seconds=elapsed,
                    html_path="",
                )
                for product in batch
            ]
            return rows, page
        elapsed = round(time.monotonic() - item_started, 2)
        html_file = self.run_dir / f"worten_batch_{batch_index:04d}.html"
        if content:
            html_file.write_text(content, encoding="utf-8")
        rows = self.classify_batch_from_cards(batch, cards, content or "", title or "", final_url or self.seller_search_url, elapsed, str(html_file))
        return rows, page


class ConforamaChecker(LeroyChecker):
    marketplace_key = "conforama"
    api_url = "https://api.empathy.co/search/v1/query/conforama/skusearch"

    def __init__(self, products: list[InputProduct], *args, **kwargs):
        self.reference_map = load_conforama_reference_map()
        mapped_products = [self.with_conforama_reference(product) for product in products]
        super().__init__(mapped_products, *args, **kwargs)
        self.active_worker_configs = self.active_worker_configs[:1]
        self.worker_configs = self.active_worker_configs
        self.worker_count = len(self.active_worker_configs)
        self.conforama_request_lock = threading.Lock()
        self.conforama_response_cache: dict[str, dict] = {}
        self.conforama_next_request_at = 0.0
        self.conforama_blocked_until = 0.0
        self.conforama_min_interval_seconds = 5.0

    def with_conforama_reference(self, product: InputProduct) -> InputProduct:
        reference = (product.reference or "").strip().upper()
        if not reference.startswith("MKP"):
            reference = self.reference_map.get(product.ean.strip(), "")
        return InputProduct(
            ean=product.ean,
            reference=reference,
            quantity=product.quantity,
            price=product.price,
        )

    def build_api_url(self, reference: str) -> str:
        params = urllib.parse.urlencode(
            {
                "internal": "true",
                "query": reference,
                "origin": "url:external",
                "start": "0",
                "rows": "10",
                "instance": "conforama",
                "lang": "es",
                "store": "es",
                "scope": "desktop",
                "currency": "EUR",
            }
        )
        return f"{self.api_url}?{params}"

    def format_price(self, value) -> str:
        if value is None or value == "":
            return ""
        try:
            number = float(value)
        except (TypeError, ValueError):
            return str(value)
        return f"{number:.2f}".replace(".", ",") + " €"

    def fetch_reference(self, reference: str) -> dict:
        reference = reference.strip().upper()
        with self.conforama_request_lock:
            cached = self.conforama_response_cache.get(reference)
            if cached is not None:
                return cached

            now = time.monotonic()
            blocked_wait = self.conforama_blocked_until - now
            if blocked_wait > 0:
                raise ConforamaApiBlocked(0, int(blocked_wait))
            wait_seconds = self.conforama_next_request_at - now
            if wait_seconds > 0:
                time.sleep(wait_seconds)
            self.conforama_next_request_at = time.monotonic() + self.conforama_min_interval_seconds

        request = urllib.request.Request(
            self.build_api_url(reference),
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/126.0.0.0 Safari/537.36"
                ),
                "Accept": "application/json, text/plain, */*",
                "Accept-Language": "es-ES,es;q=0.9",
                "Origin": "https://www.conforama.es",
                "Referer": build_marketplace_search_url(self.marketplace_key, [reference]),
                "Sec-Fetch-Site": "cross-site",
                "Sec-Fetch-Mode": "cors",
                "Sec-Fetch-Dest": "empty",
                "x-origin": "https://www.conforama.es",
            },
        )
        try:
            with urllib.request.urlopen(request, timeout=max(self.retry_timeout_ms / 1000, 10)) as response:
                payload = response.read().decode("utf-8", errors="replace")
        except urllib.error.HTTPError as exc:
            retry_after = 0
            try:
                retry_after = int(exc.headers.get("Retry-After", "0") or 0)
            except (TypeError, ValueError):
                retry_after = 0
            if exc.code in {403, 429, 503}:
                backoff = max(retry_after, 90)
                with self.conforama_request_lock:
                    self.conforama_blocked_until = max(self.conforama_blocked_until, time.monotonic() + backoff)
                raise ConforamaApiBlocked(exc.code, retry_after) from exc
            raise
        data = json.loads(payload)
        with self.conforama_request_lock:
            self.conforama_response_cache[reference] = data
        return data

    def check_conforama_product(self, product: InputProduct, index: int) -> DashboardResult:
        started = time.monotonic()
        reference = product.reference.strip().upper()
        search_url = build_marketplace_search_url(self.marketplace_key, [reference or product.ean])
        if not reference.startswith("MKP"):
            return DashboardResult(
                ean=product.ean,
                reference=product.reference,
                feed_quantity=product.quantity,
                feed_price=product.price,
                light="gray",
                status="INCIERTA",
                final_url=search_url,
                title="",
                sku="",
                seller_name="",
                marketplace_price="",
                product_availability="unknown",
                offer_available=False,
                buybox_ok="unknown",
                challenge_detected=False,
                reason="sin referencia MKP local para buscar en Conforama",
                elapsed_seconds=round(time.monotonic() - started, 2),
                html_path="",
            )

        try:
            data = self.fetch_reference(reference)
        except ConforamaApiBlocked as exc:
            self.circuit_open = True
            self.stop_requested = True
            return DashboardResult(
                ean=product.ean,
                reference=reference,
                feed_quantity=product.quantity,
                feed_price=product.price,
                light="gray",
                status="INCIERTA",
                final_url=search_url,
                title="",
                sku=reference,
                seller_name="",
                marketplace_price="",
                product_availability="unknown",
                offer_available=False,
                buybox_ok="unknown",
                challenge_detected=True,
                reason=str(exc),
                elapsed_seconds=round(time.monotonic() - started, 2),
                html_path="",
            )
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError, OSError) as exc:
            return DashboardResult(
                ean=product.ean,
                reference=reference,
                feed_quantity=product.quantity,
                feed_price=product.price,
                light="gray",
                status="INCIERTA",
                final_url=search_url,
                title="",
                sku=reference,
                seller_name="",
                marketplace_price="",
                product_availability="unknown",
                offer_available=False,
                buybox_ok="unknown",
                challenge_detected=False,
                reason=f"error API Conforama: {exc}",
                elapsed_seconds=round(time.monotonic() - started, 2),
                html_path="",
            )

        catalog = data.get("catalog", {}) if isinstance(data, dict) else {}
        content = catalog.get("content", []) if isinstance(catalog, dict) else []
        if not isinstance(catalog, dict) or not isinstance(content, list) or "numFound" not in catalog:
            json_file = self.run_dir / f"{index:04d}_{product.ean}_conforama_suspicious.json"
            json_file.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            return DashboardResult(
                ean=product.ean,
                reference=reference,
                feed_quantity=product.quantity,
                feed_price=product.price,
                light="gray",
                status="INCIERTA",
                final_url=search_url,
                title="",
                sku=reference,
                seller_name="",
                marketplace_price="",
                product_availability="unknown",
                offer_available=False,
                buybox_ok="unknown",
                challenge_detected=False,
                reason="respuesta API Conforama incompleta o sospechosa",
                elapsed_seconds=round(time.monotonic() - started, 2),
                html_path=str(json_file),
            )
        exact = None
        for item in content:
            if isinstance(item, dict) and str(item.get("id", "")).strip().upper() == reference:
                exact = item
                break

        if not exact:
            json_file = self.run_dir / f"{index:04d}_{product.ean}_conforama.json"
            json_file.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            return DashboardResult(
                ean=product.ean,
                reference=reference,
                feed_quantity=product.quantity,
                feed_price=product.price,
                light="red",
                status="NO_VIVA",
                final_url=search_url,
                title="",
                sku=reference,
                seller_name="",
                marketplace_price="",
                product_availability="not_found",
                offer_available=False,
                buybox_ok="n/a",
                challenge_detected=False,
                reason="MKP no encontrado en Conforama",
                elapsed_seconds=round(time.monotonic() - started, 2),
                html_path=str(json_file),
            )

        price = exact.get("currentPrice", exact.get("salePrice", exact.get("price", "")))
        return DashboardResult(
            ean=product.ean,
            reference=reference,
            feed_quantity=product.quantity,
            feed_price=product.price,
            light="green",
            status="OK",
            final_url=str(exact.get("url") or search_url),
            title=str(exact.get("name") or ""),
            sku=reference,
            seller_name="",
            marketplace_price=self.format_price(price),
            product_availability="available",
            offer_available=True,
            buybox_ok="n/a",
            challenge_detected=False,
            reason="MKP exacto visible en Conforama",
            elapsed_seconds=round(time.monotonic() - started, 2),
            html_path="",
        )

    def browser_worker(self, worker_id: int, work_queue: queue.Queue, events: queue.Queue) -> None:
        final_status = "cerrado"
        try:
            events.put(("worker_status", worker_id, "activo"))
            total = len(self.products)
            while not self.stop_requested:
                try:
                    index, product = work_queue.get_nowait()
                except queue.Empty:
                    break
                try:
                    row = self.check_conforama_product(product, index)
                    completed = self.record_result(row, index, total, events)
                    events.put(("result", completed, total, row))
                    if not self.stop_requested and completed < total:
                        time.sleep(self.delay)
                finally:
                    work_queue.task_done()
        except Exception as exc:
            self.stop_requested = True
            events.put(("error", f"worker {worker_id}: {exc}"))
        finally:
            events.put(("worker_status", worker_id, final_status))

    def run(self, events: queue.Queue) -> None:
        self.events = events
        started = time.monotonic()
        self.run_dir.mkdir(parents=True, exist_ok=True)
        events.put(("run_dir", str(self.run_dir)))
        events.put(("phase3", self.worker_count, self.circuit_breaker_threshold))
        for config in self.worker_configs:
            events.put(
                (
                    "worker_config",
                    config.id,
                    config.enabled,
                    config.proxy_configured,
                    config.user_agent_configured,
                )
            )
            if not config.enabled:
                events.put(("worker_status", config.id, "desactivado"))

        work_queue: queue.Queue = queue.Queue()
        for index, product in enumerate(self.products, start=1):
            work_queue.put((index, product))

        worker_threads = []
        for config in self.active_worker_configs:
            thread = threading.Thread(target=self.browser_worker, args=(config.id, work_queue, events), daemon=True)
            worker_threads.append(thread)
            thread.start()

        for thread in worker_threads:
            thread.join()

        elapsed = round(time.monotonic() - started, 1)
        summary = str(self.summary_path)
        if self.circuit_open:
            summary = f"{summary} - corte tecnico activo"
        events.put(("done", elapsed, summary))


def completed_count_safe_index(row: DashboardResult, products: list[InputProduct]) -> int:
    for index, product in enumerate(products, start=1):
        if product.ean == row.ean:
            return index
    return 0


def create_checker(marketplace_key: str, *args, **kwargs):
    if marketplace_key == "carrefour":
        return CarrefourChecker(*args, **kwargs)
    if marketplace_key == "worten":
        return WortenChecker(*args, **kwargs)
    if marketplace_key == "conforama":
        return ConforamaChecker(*args, **kwargs)
    return LeroyChecker(*args, **kwargs)


class MarketplaceDialog(tk.Toplevel):
    def __init__(self, master: tk.Tk):
        super().__init__(master)
        self.title("Marketplace")
        self.resizable(False, False)
        self.result = "Leroy Merlin"
        self.configure(bg="#f6f7f4")
        self.transient(master)
        self.grab_set()

        frame = ttk.Frame(self, padding=22)
        frame.grid(row=0, column=0, sticky="nsew")
        ttk.Label(frame, text="Marketplace a checkear", font=("Segoe UI", 13, "bold")).grid(row=0, column=0, sticky="w")
        self.combo = ttk.Combobox(frame, state="readonly", values=["Leroy Merlin"], width=30)
        self.combo.set("Leroy Merlin")
        self.combo.grid(row=1, column=0, pady=(12, 18), sticky="ew")
        ttk.Button(frame, text="Continuar", command=self.accept).grid(row=2, column=0, sticky="e")
        self.bind("<Return>", lambda _event: self.accept())
        self.wait_visibility()
        self.focus_force()

    def accept(self) -> None:
        self.result = self.combo.get()
        self.destroy()


class DashboardApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Offer Watcher")
        self.geometry("1240x760")
        self.minsize(1100, 680)
        self.configure(bg="#f5f6f2")
        self.events: queue.Queue = queue.Queue()
        self.products: list[InputProduct] = []
        self.checker: LeroyChecker | None = None
        self.worker: threading.Thread | None = None
        self.counts = {"green": 0, "yellow": 0, "red": 0, "gray": 0}

        self.style = ttk.Style(self)
        self.style.theme_use("clam")
        self.style.configure("TFrame", background="#f5f6f2")
        self.style.configure("TLabel", background="#f5f6f2", foreground="#1f2a24", font=("Segoe UI", 10))
        self.style.configure("Header.TLabel", font=("Segoe UI", 16, "bold"))
        self.style.configure("TButton", font=("Segoe UI", 10), padding=8)
        self.style.configure("Treeview", rowheight=30, font=("Segoe UI", 9), fieldbackground="#ffffff")
        self.style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"))

        dialog = MarketplaceDialog(self)
        self.marketplace = dialog.result
        self.build_ui()
        self.after(200, self.process_events)

    def build_ui(self) -> None:
        top = ttk.Frame(self, padding=(20, 16, 20, 8))
        top.pack(fill="x")

        logo = tk.Canvas(top, width=210, height=58, bg="#f5f6f2", highlightthickness=0)
        logo.grid(row=0, column=0, rowspan=2, sticky="w")
        logo.create_rectangle(0, 6, 190, 54, fill="#007a3d", outline="#007a3d")
        logo.create_polygon(18, 43, 48, 15, 78, 43, fill="white", outline="white")
        logo.create_text(126, 30, text="LEROY\nMERLIN", fill="white", font=("Segoe UI", 12, "bold"), justify="center")

        ttk.Label(top, text="Disponibilidad Marketplace", style="Header.TLabel").grid(row=0, column=1, sticky="w")
        self.marketplace_var = tk.StringVar(value=self.marketplace)
        ttk.Combobox(top, textvariable=self.marketplace_var, values=["Leroy Merlin"], state="readonly", width=24).grid(
            row=1, column=1, sticky="w", pady=(6, 0)
        )

        status_frame = ttk.Frame(top)
        status_frame.grid(row=0, column=2, rowspan=2, sticky="e")
        top.columnconfigure(2, weight=1)
        self.green_label = ttk.Label(status_frame, text="● 0 OK", foreground="#11823b", font=("Segoe UI", 11, "bold"))
        self.yellow_label = ttk.Label(status_frame, text="● 0 Revisar", foreground="#b77900", font=("Segoe UI", 11, "bold"))
        self.red_label = ttk.Label(status_frame, text="● 0 Fuera", foreground="#bd1e24", font=("Segoe UI", 11, "bold"))
        self.gray_label = ttk.Label(status_frame, text="● 0 Tecnico", foreground="#6b7280", font=("Segoe UI", 11, "bold"))
        self.green_label.grid(row=0, column=0, padx=12)
        self.yellow_label.grid(row=0, column=1, padx=12)
        self.red_label.grid(row=0, column=2, padx=12)
        self.gray_label.grid(row=0, column=3, padx=12)

        controls = ttk.Frame(self, padding=(20, 8, 20, 8))
        controls.pack(fill="x")
        ttk.Label(controls, text="EAN").grid(row=0, column=0, sticky="w")
        self.ean_var = tk.StringVar()
        ttk.Entry(controls, textvariable=self.ean_var, width=32).grid(row=1, column=0, sticky="ew", padx=(0, 12))
        ttk.Button(controls, text="Agregar EAN", command=self.add_single_ean).grid(row=1, column=1, padx=(0, 8))
        ttk.Button(controls, text="Cargar CSV bulk", command=self.load_bulk).grid(row=1, column=2, padx=(0, 8))
        ttk.Button(controls, text="Limpiar lista", command=self.clear_list).grid(row=1, column=3, padx=(0, 8))
        ttk.Button(controls, text="Limpiar perfil", command=self.clean_profile).grid(row=1, column=4, padx=(0, 8))
        controls.columnconfigure(0, weight=1)

        options = ttk.Frame(self, padding=(20, 0, 20, 8))
        options.pack(fill="x")
        self.expected_seller_var = tk.StringVar(value="NEWLUX GROUP")
        self.delay_var = tk.StringVar(value="2")
        ttk.Label(options, text="Seller esperado").grid(row=0, column=0, sticky="w")
        ttk.Entry(options, textvariable=self.expected_seller_var, width=28).grid(row=0, column=1, sticky="w", padx=(8, 18))
        ttk.Label(options, text="Delay").grid(row=0, column=2, sticky="w")
        ttk.Entry(options, textvariable=self.delay_var, width=8).grid(row=0, column=3, sticky="w", padx=(8, 18))
        ttk.Button(options, text="Analizar", command=self.start_analysis).grid(row=0, column=4, padx=(0, 8))
        ttk.Button(options, text="Parar", command=self.stop_analysis).grid(row=0, column=5, padx=(0, 8))
        self.progress_var = tk.StringVar(value="Sin ejecucion")
        ttk.Label(options, textvariable=self.progress_var).grid(row=0, column=6, sticky="w", padx=(18, 0))

        table_frame = ttk.Frame(self, padding=(20, 8, 20, 20))
        table_frame.pack(fill="both", expand=True)
        columns = (
            "light",
            "ean",
            "reference",
            "stock",
            "feed_price",
            "status",
            "seller",
            "market_price",
            "availability",
            "reason",
            "url",
        )
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings")
        headings = {
            "light": "",
            "ean": "EAN",
            "reference": "SKU/ref",
            "stock": "Stock",
            "feed_price": "Precio feed",
            "status": "Estado",
            "seller": "Seller",
            "market_price": "Precio LM",
            "availability": "Disponibilidad",
            "reason": "Motivo",
            "url": "URL",
        }
        widths = {
            "light": 46,
            "ean": 122,
            "reference": 110,
            "stock": 70,
            "feed_price": 86,
            "status": 150,
            "seller": 150,
            "market_price": 86,
            "availability": 180,
            "reason": 260,
            "url": 380,
        }
        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(col, width=widths[col], minwidth=40, anchor="w")
        self.tree.tag_configure("green", background="#eaf6ee", foreground="#0d5f2b")
        self.tree.tag_configure("yellow", background="#fff7df", foreground="#6e5200")
        self.tree.tag_configure("red", background="#fdecec", foreground="#8f171d")
        self.tree.tag_configure("gray", background="#f1f3f5", foreground="#4b5563")
        self.tree.pack(side="left", fill="both", expand=True)
        scroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")

    def add_single_ean(self) -> None:
        ean = self.ean_var.get().strip()
        if not ean:
            return
        self.products.append(InputProduct(ean=ean))
        self.tree.insert("", "end", values=("○", ean, "", "", "", "PENDIENTE", "", "", "", "", ""), tags=("pending",))
        self.ean_var.set("")
        self.progress_var.set(f"{len(self.products)} productos cargados")

    def load_bulk(self) -> None:
        path = filedialog.askopenfilename(
            title="Selecciona CSV",
            filetypes=[("CSV", "*.csv"), ("Texto", "*.txt"), ("Todos", "*.*")],
        )
        if not path:
            return
        try:
            self.products = read_products_csv(Path(path))
        except Exception as exc:
            messagebox.showerror("Error CSV", str(exc))
            return
        self.tree.delete(*self.tree.get_children())
        for product in self.products:
            self.tree.insert(
                "",
                "end",
                values=("○", product.ean, product.reference, product.quantity, product.price, "PENDIENTE", "", "", "", "", ""),
            )
        self.progress_var.set(f"{len(self.products)} productos cargados")

    def clear_list(self) -> None:
        if self.worker and self.worker.is_alive():
            messagebox.showwarning("En ejecucion", "Para el analisis antes de limpiar la lista.")
            return
        self.products = []
        self.tree.delete(*self.tree.get_children())
        self.counts = {"green": 0, "yellow": 0, "red": 0, "gray": 0}
        self.update_counts()
        self.progress_var.set("Lista limpia")

    def clean_profile(self) -> None:
        if self.worker and self.worker.is_alive():
            messagebox.showwarning("En ejecucion", "Para el analisis antes de limpiar el perfil.")
            return
        if PROFILE_DIR.exists():
            shutil.rmtree(PROFILE_DIR)
        self.progress_var.set("Perfil dedicado limpiado")

    def start_analysis(self) -> None:
        if not self.products:
            self.add_single_ean()
        if not self.products:
            messagebox.showwarning("Sin productos", "Agrega un EAN o carga un CSV.")
            return
        if self.worker and self.worker.is_alive():
            return
        try:
            delay = float(self.delay_var.get().replace(",", "."))
        except ValueError:
            delay = 2.0
        self.counts = {"green": 0, "yellow": 0, "red": 0, "gray": 0}
        self.update_counts()
        self.tree.delete(*self.tree.get_children())
        for product in self.products:
            self.tree.insert("", "end", values=("○", product.ean, product.reference, product.quantity, product.price, "PENDIENTE", "", "", "", "", ""))
        self.checker = LeroyChecker(self.products, self.expected_seller_var.get(), delay)
        self.worker = threading.Thread(target=self.checker.run, args=(self.events,), daemon=True)
        self.worker.start()
        self.progress_var.set("Analizando...")

    def stop_analysis(self) -> None:
        if self.checker:
            self.checker.stop()
            self.progress_var.set("Parando al finalizar el producto actual...")

    def update_counts(self) -> None:
        self.green_label.configure(text=f"● {self.counts['green']} OK")
        self.yellow_label.configure(text=f"● {self.counts['yellow']} Revisar")
        self.red_label.configure(text=f"● {self.counts['red']} Fuera")
        self.gray_label.configure(text=f"● {self.counts['gray']} Tecnico")

    def process_events(self) -> None:
        try:
            while True:
                event = self.events.get_nowait()
                kind = event[0]
                if kind == "result":
                    _, index, total, row = event
                    children = self.tree.get_children()
                    item = children[index - 1] if index - 1 < len(children) else self.tree.insert("", "end")
                    symbol = {"green": "●", "yellow": "●", "red": "●"}.get(row.light, "●")
                    self.tree.item(
                        item,
                        values=(
                            symbol,
                            row.ean,
                            row.reference,
                            row.feed_quantity,
                            row.feed_price,
                            row.status,
                            row.seller_name,
                            row.marketplace_price,
                            row.product_availability,
                            row.reason,
                            row.final_url,
                        ),
                        tags=(row.light,),
                    )
                    if row.light in self.counts:
                        self.counts[row.light] += 1
                        self.update_counts()
                    self.progress_var.set(f"{index}/{total} analizados")
                elif kind == "done":
                    _, elapsed, summary = event
                    self.progress_var.set(f"Terminado en {elapsed}s · {summary}")
                elif kind == "run_dir":
                    self.progress_var.set(f"Salida: {event[1]}")
                elif kind == "error":
                    messagebox.showerror("Error", event[1])
        except queue.Empty:
            pass
        self.after(200, self.process_events)


def main() -> None:
    app = DashboardApp()
    app.mainloop()


if __name__ == "__main__":
    main()
